"""Score a finished proposal deck the way the client's evaluation panel would - on the request's own criteria, then on logic,
consistency, integrity, faithfulness to the base template and visual quality - with an evaluator from another model family
than the page authors and the reviewer.

    python hosts/responses_api/evaluate_proposal.py <project> [--deck exports/X.pptx] [--criteria sources/evaluation-criteria.md]
        [--model claude-opus-5-5] [--effort medium]

The panel sees what a client sees: the request's documents, the exported deck as PowerPoint draws it (one image per slide) and the
deck's text. For the integrity check it also sees the bidder's own knowledge base, so it can tell a stated fact from an invented
one. For style it sees the official template the deck was meant to follow. Alongside the judgement, the fonts and colours the
PowerPoint file actually uses are counted against the template's theme - a measurement, not an opinion.

Writes `<project>/.review/proposal_evaluation.md` and `.json`.
"""

from __future__ import annotations

import argparse
import base64
import io
import json
import re
import subprocess
import sys
import time
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
SCRIPTS = ROOT / "skills" / "ppt-master" / "scripts"
sys.path.insert(0, str(HERE))

PANEL = """You are the client's evaluation panel scoring one bidder's proposal deck. You are experienced, fair and hard to impress: you have read many proposals for this kind of work, and you score only what the deck shows. You do not reward length, adjectives or the bidder's claims about itself; you reward a specific, credible, well-reasoned answer to THIS request.

You receive: the request's documents (what the client asked for and how it scores), the deck as PowerPoint renders it (one image per slide, in order) with each slide's text, the bidder's own knowledge base (for the integrity check only - it tells you what the bidder can actually stand behind), and the official template the deck was meant to be built on.

Work in this order and be concrete - cite slide numbers for every finding.
1. CRITERIA. For each evaluation criterion the request states, in its order and with its weight: what the deck offers against it, where (slides), what is missing or weak, and a score on the request's own scale if it gives one (say which scale), else 0-10. If the request defines no criteria, derive the four to six a panel for this work would use and say so.
2. LOGIC. Does the deck argue - a governing thought, a storyline that builds, recommendations that follow from a stated understanding of the client's problem, a solution that answers the stated scope? Score 0-10.
3. CONSISTENCY. Numbers, names, dates and durations that must agree across slides (price totals and their breakdowns, team days, plan weeks, milestone dates, named components). Recompute any total you can from its parts. List every disagreement. Score 0-10.
4. INTEGRITY. Anything stated as fact that neither the request nor the bidder's knowledge base supports (named people, clients, credentials, measured results, certifications) is an invention - list each. Placeholders marked `[To be provided: ...]` are honest and are not penalised; say how many there are and whether they sit where a reader expects them. Score 0-10.
5. TEMPLATE FIDELITY. Against the official template's images: are its layouts, marks, colours, fonts, title placement and footer carried through every slide? Name slides that depart and how. Score 0-10.
6. VISUAL QUALITY. As a consulting page: hierarchy, alignment, density, legibility, figures that carry meaning, nothing broken (overlaps, clipped text, stray elements). Name the three weakest slides and why. Score 0-10.
7. SELF-CONTAINMENT. The deck is sent to a senior reader with no context and nobody to ask. Name every term, abbreviation, code, internal name or shorthand used without explanation, every page whose point needs a verbal rescue, and whether the deck has a clear beginning, middle and end. Score 0-10.
8. VERDICT. Would this deck make the shortlist? The three changes that would raise the score most.

End with ONE fenced ```json block, exactly this shape:
{"criteria": [{"name": "...", "weight": "... as the request states it, or null", "scale": "...", "score": 0, "max": 10, "slides": [1], "gaps": "..."}],
 "weighted_score_pct": 0, "logic": 0, "consistency": 0, "integrity": 0, "template_fidelity": 0, "visual_quality": 0, "self_containment": 0, "unexplained_terms": ["..."],
 "inventions": ["..."], "inconsistencies": ["..."], "placeholders": 0, "shortlist": true, "top_changes": ["...", "...", "..."]}
weighted_score_pct is the criteria scores weighted by the request's weights (or equally if it gives none), as a percentage of the maximum."""


def slide_texts(deck: Path) -> list[str]:
    from pptx import Presentation
    texts = []
    for slide in Presentation(str(deck)).slides:
        parts = []
        shapes = sorted(slide.shapes, key=lambda sh: ((sh.top or 0) // 20, sh.left or 0))
        for shape in shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        texts.append("\n".join(parts))
    return texts


def style_measure(deck: Path, manifest: Path | None) -> dict:
    """Fonts and colours the slides actually use, counted against the template's theme."""
    fonts: dict[str, int] = {}
    colours: dict[str, int] = {}
    with zipfile.ZipFile(deck) as z:
        for name in z.namelist():
            if re.match(r"ppt/slides/slide\d+\.xml$", name):
                xml = z.read(name).decode("utf-8", errors="replace")
                for face in re.findall(r"<a:latin typeface=\"([^\"]+)\"", xml):
                    fonts[face] = fonts.get(face, 0) + 1
                for hex_ in re.findall(r"<a:srgbClr val=\"([0-9A-Fa-f]{6})\"", xml):
                    colours[hex_.upper()] = colours.get(hex_.upper(), 0) + 1
    result = {"fonts": fonts, "colours": dict(sorted(colours.items(), key=lambda kv: -kv[1])[:30])}
    if manifest and manifest.is_file():
        theme = json.loads(manifest.read_text(encoding="utf-8")).get("theme") or {}
        palette = {v.lstrip("#").upper() for v in (theme.get("colors") or {}).values() if isinstance(v, str)}
        theme_fonts = {v for v in (theme.get("fonts") or {}).values() if isinstance(v, str)}

        def near(hex_: str) -> bool:
            r, g, b = (int(hex_[i:i + 2], 16) for i in (0, 2, 4))
            return any(abs(r - int(p[0:2], 16)) + abs(g - int(p[2:4], 16)) + abs(b - int(p[4:6], 16)) <= 24 for p in palette) or hex_ in {"FFFFFF", "000000"}

        total_c = sum(colours.values()) or 1
        total_f = sum(fonts.values()) or 1
        result.update({"theme_palette": sorted(palette), "theme_fonts": sorted(theme_fonts),
                       "colour_uses_on_theme_pct": round(100 * sum(n for c, n in colours.items() if near(c)) / total_c, 1),
                       "font_uses_on_theme_pct": round(100 * sum(n for f, n in fonts.items() if f.startswith("+") or f in theme_fonts or any(f.startswith(t) for t in theme_fonts)) / total_f, 1),  # "+mj-lt"/"+mn-lt" ARE the theme's fonts
                       "off_theme_colours": [c for c in colours if not near(c)][:15]})
    return result


def image_block(path: Path, width: int = 1280) -> dict:
    from PIL import Image
    with Image.open(path) as image:
        image = image.convert("RGB")
        if image.width > width:
            image = image.resize((width, int(image.height * width / image.width)), Image.LANCZOS)
        buffer = io.BytesIO()
        image.save(buffer, format="JPEG", quality=85)
    return {"type": "image", "source": {"type": "base64", "media_type": "image/jpeg", "data": base64.b64encode(buffer.getvalue()).decode("ascii")}}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("project")
    parser.add_argument("--deck")
    parser.add_argument("--criteria", help="default: sources/evaluation-criteria.md when present")
    parser.add_argument("--model", default="claude-opus-5-5")
    parser.add_argument("--effort", default="medium")
    parser.add_argument("--panel-file", help="the evaluator's instructions for a deck that is not a proposal (e.g. a run brief's rubric); must end by asking for a fenced json block")
    parser.add_argument("--out-name", default="proposal_evaluation", help="file stem under .review/")
    args = parser.parse_args()
    panel = Path(args.panel_file).read_text(encoding="utf-8") if args.panel_file else PANEL
    import anthropic_backend
    from host import _key

    project = Path(args.project).resolve()
    deck = Path(args.deck).resolve() if args.deck else max((project / "exports").glob("*.pptx"), key=lambda p: p.stat().st_mtime)
    render = deck.with_suffix(".render")
    if not sorted(render.glob("slide-*.png")):
        subprocess.run([sys.executable, str(SCRIPTS / "pptx_render.py"), str(deck), "--out", str(render)], capture_output=True, text=True)
    slides = sorted(render.glob("slide-*.png"))
    if not slides:
        raise SystemExit("no PowerPoint render of the deck: the panel judges what PowerPoint shows")
    texts = slide_texts(deck)
    sources = project / "sources"
    kb = [p for p in sources.glob("*") if p.is_file() and re.search(r"knowledge|firm", p.name, re.I)]
    request_docs = [p for p in sorted(sources.glob("*")) if p.is_file() and p.suffix.lower() in {".md", ".txt"} and p not in kb and p.name != "base-template.md"]
    criteria = Path(args.criteria).resolve() if args.criteria else sources / "evaluation-criteria.md"
    manifest = project / "base_template" / "import" / "analysis" / "manifest.json"
    measured = style_measure(deck, manifest if manifest.is_file() else None)

    content: list[dict] = [{"type": "text", "text": "THE REQUEST'S DOCUMENTS\n\n" + "\n\n".join(f"<document name=\"{p.name}\">\n{p.read_text(encoding='utf-8', errors='replace')}\n</document>" for p in request_docs)}]
    if criteria.is_file() and criteria not in request_docs:
        content.append({"type": "text", "text": f"THE EVALUATION CRITERIA\n\n{criteria.read_text(encoding='utf-8', errors='replace')}"})
    content.append({"type": "text", "text": "THE BIDDER'S OWN KNOWLEDGE BASE (for the integrity check only)\n\n" + "\n\n".join(p.read_text(encoding="utf-8", errors="replace") for p in kb)})
    template_images = sorted((project / "base_template" / "layouts").glob("layout_*.png"))[:4] + sorted((project / "base_template" / "render").glob("slide-*.png"))[:4]
    if template_images:
        content.append({"type": "text", "text": "THE OFFICIAL TEMPLATE the deck was meant to follow (its layouts, then its sample slides):"})
        content += [image_block(p, 960) for p in template_images]
        content.append({"type": "text", "text": "MEASURED in the PowerPoint file (fonts and colours actually used, against the template's theme):\n" + json.dumps(measured, indent=1)})
    content.append({"type": "text", "text": f"THE DECK - {len(slides)} slides as PowerPoint renders them, each followed by its text:"})
    for number, png in enumerate(slides, start=1):
        content.append({"type": "text", "text": f"--- Slide {number} ---"})
        content.append(image_block(png))
        content.append({"type": "text", "text": texts[number - 1] if number <= len(texts) else ""})

    request = {"model": args.model, "max_tokens": 128000, "system": panel, "messages": [{"role": "user", "content": content}],
               "output_config": {"effort": args.effort}}
    started = time.time()
    message = anthropic_backend._stream_with_retry(_key("ANTHROPIC_API_KEY"), request)
    seconds = time.time() - started
    text = "\n".join(b.text for b in message.content if getattr(b, "type", "") == "text")
    usage = message.usage
    price = anthropic_backend.PRICES.get(args.model) or {}
    cost = ((usage.input_tokens or 0) * price.get("input", 0) + (usage.output_tokens or 0) * price.get("output", 0)
            + (getattr(usage, "cache_read_input_tokens", 0) or 0) * price.get("cache_read", 0) + (getattr(usage, "cache_creation_input_tokens", 0) or 0) * price.get("cache_write", 0))
    match = re.search(r"```json\s*(\{.*?\})\s*```", text, re.S)
    scores = None
    if match:
        try:
            scores = json.loads(match.group(1))
        except json.JSONDecodeError:
            scores = None
    out = project / ".review"
    out.mkdir(parents=True, exist_ok=True)
    header = (f"# Proposal evaluation - {deck.name}\n\nEvaluator {args.model} @ {args.effort}; {seconds:.0f} s; USD {cost:.2f}; "
              f"{usage.input_tokens} input / {usage.output_tokens} output tokens.\n\n## Measured style\n\n```json\n{json.dumps(measured, indent=1)}\n```\n\n## Panel\n\n")
    (out / f"{args.out_name}.md").write_text(header + text.strip() + "\n", encoding="utf-8")
    (out / f"{args.out_name}.json").write_text(json.dumps({"deck": str(deck), "model": args.model, "effort": args.effort, "seconds": round(seconds, 1),
                                                              "cost_usd": round(cost, 4), "measured": measured, "scores": scores}, indent=1), encoding="utf-8")
    print(f"evaluation: {out / (args.out_name + '.md')} ({seconds:.0f} s, USD {cost:.2f})")
    if scores:
        print(json.dumps({k: scores.get(k) for k in ("weighted_score_pct", "logic", "consistency", "integrity", "template_fidelity", "visual_quality", "self_containment", "shortlist")}))
    return 0


if __name__ == "__main__":
    sys.exit(main())
