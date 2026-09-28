"""Fork tools of the consulting-quality profile: the review journal and reference matching."""

import json
import os
import sys
from pathlib import Path

os.environ["PPT_MASTER_NO_LINT"] = "1"  # the geometry lint needs a browser; its rules are tested below on plain geometry
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import page_lint  # noqa: E402
import page_review  # noqa: E402
import reference_library  # noqa: E402


def _project(tmp_path: Path) -> Path:
    (tmp_path / "svg_output").mkdir()
    (tmp_path / "svg_output" / "01_cover.svg").write_text("<svg/>", encoding="utf-8")
    return tmp_path


def _render_args(project: Path, **extra):
    import argparse
    return argparse.Namespace(project=str(project), page="01_cover", crop=None, **extra)


def test_revision_counts_changed_content_only(tmp_path, monkeypatch):
    project = _project(tmp_path)
    png = project / "p.png"
    png.write_bytes(b"x")
    monkeypatch.setattr(page_review, "_render_and_lint", lambda root, svg, want_lint: ([{"ok": True, "path": str(png)}], None, None))
    page_review.cmd_render(_render_args(project))
    page_review.cmd_render(_render_args(project))  # same content: a second look, not a revision
    entry = page_review.load_journal(project)["pages"]["01_cover"]
    assert (entry["renders"], entry["revisions"], len(entry["backups"])) == (2, 0, 1)
    (project / "svg_output" / "01_cover.svg").write_text("<svg id='b'/>", encoding="utf-8")
    page_review.cmd_render(_render_args(project))
    entry = page_review.load_journal(project)["pages"]["01_cover"]
    assert (entry["revisions"], len(entry["backups"])) == (1, 2)


def test_restore_brings_back_an_earlier_revision(tmp_path, monkeypatch):
    project = _project(tmp_path)
    png = project / "p.png"
    png.write_bytes(b"x")
    monkeypatch.setattr(page_review, "_render_and_lint", lambda root, svg, want_lint: ([{"ok": True, "path": str(png)}], None, None))
    page_review.cmd_render(_render_args(project))
    svg = project / "svg_output" / "01_cover.svg"
    svg.write_text("<svg id='worse'/>", encoding="utf-8")
    page_review.cmd_render(_render_args(project))
    import argparse
    page_review.cmd_restore(argparse.Namespace(project=str(project), page="01_cover", rev=0))
    assert svg.read_text(encoding="utf-8") == "<svg/>"


def test_outcome_needs_a_look_first(tmp_path):
    import argparse
    import pytest
    project = _project(tmp_path)
    with pytest.raises(SystemExit):
        page_review.cmd_note(argparse.Namespace(project=str(project), page="01_cover", outcome="accepted", text="fine"))


def test_rejected_entries_never_match_positively():
    good = {"id": "a", "labels": {"communication_form": "timeline", "devices": ["phase ruler"]}}
    score_good, _ = reference_library.score(good, "timeline", {"phase", "ruler"}, None)
    assert score_good > reference_library.WEAK_SCORE
    assert reference_library._rejected({"verdict": "rejected"}) and not reference_library._rejected(good)


def test_form_outranks_shared_words():
    same_form = {"labels": {"communication_form": "architecture"}}
    other_form = {"labels": {"communication_form": "cover", "devices": ["hub", "band"]}}
    assert reference_library.score(same_form, "architecture", {"hub", "band"}, None)[0] > \
        reference_library.score(other_form, "architecture", {"hub", "band"}, None)[0]


def _deck(path: Path, pages: int = 3) -> Path:
    import fitz
    doc = fitz.open()
    for index in range(pages):
        page = doc.new_page(width=960, height=540)
        page.insert_text((60, 120), f"page {index + 1}", fontsize=40)
    doc.save(path)
    return path


def test_build_indexes_every_page_unlabelled_and_replaces_on_rebuild(tmp_path):
    library = tmp_path / "lib"
    deck = _deck(tmp_path / "deck.pdf")
    entries = reference_library.build(library, [deck], width=800)
    assert [e["id"].rsplit(".", 1)[1] for e in entries] == ["p1", "p2", "p3"]
    assert {reference_library._form(e) for e in entries} == {reference_library.UNLABELLED}
    assert all((library / e["image"]).is_file() for e in entries)
    entries = reference_library.build(library, [deck], width=800)  # same deck again: replaced, not duplicated
    assert len(entries) == 3 and len(reference_library.load(library)) == 3


def test_sheet_stamps_ids_and_skips_rejected_and_unlabelled_for_match(tmp_path, capsys):
    import argparse
    library = tmp_path / "lib"
    reference_library.build(library, [_deck(tmp_path / "deck.pdf")], width=800)
    entries = reference_library.load(library)
    ids = [e["id"] for e in entries]
    reference_library.cmd_label(argparse.Namespace(ids=ids[:2], form="timeline", topology=None, devices=None, verdict=None, reason=None), library, entries)
    reference_library.cmd_label(argparse.Namespace(ids=ids[1:2], form=None, topology=None, devices=None, verdict="rejected", reason="equal cards"), library, entries)
    entries = reference_library.load(library)
    capsys.readouterr()
    reference_library.cmd_sheet(argparse.Namespace(form="timeline", limit=24, columns=4, exclude=[], project=None, page=None), library, entries)
    out = capsys.readouterr().out
    assert ids[0] in out and ids[1] not in out and "IMAGE:" in out  # rejected never on a sheet
    assert (library / "sheets" / "timeline.png").is_file()
    reference_library.cmd_match(argparse.Namespace(command="match", form="timeline", need="", density=None, limit=3, exclude=[], project=None, page=None), library, entries)
    out = capsys.readouterr().out
    assert ids[0] in out and ids[2] not in out  # the unlabelled page is never a positive match


def test_sheet_paginates_past_the_limit(tmp_path, capsys):
    import argparse
    library = tmp_path / "lib"
    reference_library.build(library, [_deck(tmp_path / "deck.pdf", pages=5)], width=800)
    entries = reference_library.load(library)
    reference_library.cmd_sheet(argparse.Namespace(form=reference_library.UNLABELLED, limit=2, columns=4, exclude=[], project=None, page=None), library, entries)
    out = capsys.readouterr().out
    assert out.count("IMAGE:") == 3 and (library / "sheets" / "unlabelled-3.png").is_file()


def test_show_records_delivery_in_the_journal(tmp_path, capsys):
    import argparse
    library = tmp_path / "lib"
    reference_library.build(library, [_deck(tmp_path / "deck.pdf")], width=800)
    entries = reference_library.load(library)
    (tmp_path / "proj").mkdir()
    project = _project(tmp_path / "proj")
    reference_library.cmd_show(argparse.Namespace(ids=[entries[0]["id"]], project=str(project), page="03_architecture"), library, entries)
    assert "IMAGE:" in capsys.readouterr().out
    delivered = page_review.load_journal(project)["references_delivered"]
    assert delivered[0]["id"] == entries[0]["id"] and delivered[0]["page"] == "03_architecture"


def _reviewed_project(tmp_path, monkeypatch, verdict_text):
    import argparse
    project = _project(tmp_path)
    (project / ".preview").mkdir()
    (project / ".preview" / "01_cover.png").write_bytes(b"\x89PNG")
    (project / "design_spec.md").write_text("## IX. Content Outline\n\n#### Slide 01 - Cover\n\n- **Audience move**: frame.\n- **Relationships**: none.\n\n## X. Notes\n", encoding="utf-8")
    monkeypatch.setattr(page_review, "_render_and_lint", lambda root, svg, want_lint: ([{"ok": True, "path": str(project / ".preview" / "01_cover.png")}], None, None))
    page_review.cmd_render(_render_args(project))
    calls = []
    def fake_call(payload):
        calls.append(payload)
        return verdict_text, {"input_tokens": 10, "output_tokens": 5}
    monkeypatch.setattr(page_review, "_review_call", fake_call)
    return project, calls


def test_review_needs_the_current_render_and_records_the_verdict(tmp_path, monkeypatch, capsys):
    import argparse, pytest
    project, calls = _reviewed_project(tmp_path, monkeypatch, "1. text under line\n\nVERDICT: EXECUTION_REPAIR\nBLOCKERS: 1")
    page_review.cmd_review(argparse.Namespace(project=str(project), page="01_cover"))
    review = page_review.load_journal(project)["pages"]["01_cover"]["review"]
    assert (review["verdict"], review["blockers"], review["rev"]) == ("EXECUTION_REPAIR", 1, 0)
    assert (project / review["file"]).is_file() and "Slide 01 - Cover" in calls[0]["input"][0]["content"][0]["text"]
    assert calls[0]["instructions"] == page_review.REVIEW_INSTRUCTIONS and "previous_response_id" not in calls[0]
    (project / "svg_output" / "01_cover.svg").write_text("<svg id='edited'/>", encoding="utf-8")
    with pytest.raises(SystemExit, match="render this revision first"):
        page_review.cmd_review(argparse.Namespace(project=str(project), page="01_cover"))


def test_accepted_needs_a_pass_review_of_this_revision(tmp_path, monkeypatch):
    import argparse, pytest
    project, _ = _reviewed_project(tmp_path, monkeypatch, "none found\n\nVERDICT: EXECUTION_REPAIR\nBLOCKERS: 2")
    note = lambda **kw: page_review.cmd_note(argparse.Namespace(project=str(project), page="01_cover", text="fine", **kw))
    with pytest.raises(SystemExit, match="no independent review"):
        note(outcome="accepted")
    page_review.cmd_review(argparse.Namespace(project=str(project), page="01_cover"))
    with pytest.raises(SystemExit, match="EXECUTION_REPAIR with 2 blocker"):
        note(outcome="accepted")
    note(outcome="unresolved")  # unresolved never needs a review
    note(outcome="accepted", no_review=True)  # the explicit exception is allowed and written into the note
    assert "[accepted without an independent review" in (project / ".review" / "notes" / "01_cover.md").read_text(encoding="utf-8")
    monkeypatch.setattr(page_review, "_review_call", lambda payload: ("none found\n\nVERDICT: PASS\nBLOCKERS: 0", {}))
    (project / "svg_output" / "01_cover.svg").write_text("<svg id='fixed'/>", encoding="utf-8")
    page_review.cmd_render(_render_args(project))  # a fix is a new revision; the reviewer looks at that one
    page_review.cmd_review(argparse.Namespace(project=str(project), page="01_cover"))
    note(outcome="accepted")
    assert page_review.load_journal(project)["pages"]["01_cover"]["outcome"] == "accepted"


def test_review_parses_markdown_verdicts_and_refuses_a_second_opinion_on_an_unchanged_render(tmp_path, monkeypatch):
    import argparse, pytest
    project, _ = _reviewed_project(tmp_path, monkeypatch, "## 1. Inventory\n\nnone found\n\n## VERDICT\n\nPASS - clean cover.\n\n**BLOCKERS:** 0")
    page_review.cmd_review(argparse.Namespace(project=str(project), page="01_cover"))
    review = page_review.load_journal(project)["pages"]["01_cover"]["review"]
    assert (review["verdict"], review["blockers"]) == ("PASS", 0)
    with pytest.raises(SystemExit, match="already reviewed"):
        page_review.cmd_review(argparse.Namespace(project=str(project), page="01_cover"))
    monkeypatch.setattr(page_review, "_review_call", lambda payload: ("no closing lines at all", {}))
    (project / "svg_output" / "01_cover.svg").write_text("<svg id='v2'/>", encoding="utf-8")
    page_review.cmd_render(_render_args(project))
    page_review.cmd_review(argparse.Namespace(project=str(project), page="01_cover"))
    assert page_review.load_journal(project)["pages"]["01_cover"]["review"]["verdict"] == "UNPARSED"
    page_review.cmd_review(argparse.Namespace(project=str(project), page="01_cover"))  # an unparsed review may be asked again


SPEC = """# Spec

## I. Project Information

design system here

## IX. Content Outline

> note

### Narrative

- **Governing thought**: one loop.

### Part 1: A

#### Slide 01 - Cover

- **Role**: cover
- **Audience move**: frame.
- **Relationships**: none.
- **Title**: Sentinel for PinnacleOne
- **Core message**: a proposal.

#### Slide 02 - Executive summary

- **Role**: executive_summary
- **Author tier**: workhorse
- **Title**: Appoint us
- **Core message**: approve the programme.
- **Visual scaffold**: SECRET-SCAFFOLD-02

### Part 2: B

#### Slide 03 - Architecture in your estate

- **Role**: architecture
- **Author tier**: frontier
- **Title (binding)**: Everything runs in your tenancy
- **Core message**: in-region.

## X. Speaker Notes Requirements
"""


def _runner():
    sys.path.insert(0, str(Path(__file__).resolve().parents[4] / "hosts" / "responses_api"))
    import deck_runner
    return deck_runner


def test_pages_are_parsed_with_tier_stem_and_their_own_record_only():
    pages = _runner().parse_pages(SPEC)
    assert [(p["number"], p["stem"], p["tier"], p["role"]) for p in pages] == [
        (1, "01_cover", "workhorse", "cover"), (2, "02_executive_summary", "workhorse", "executive_summary"),
        (3, "03_architecture_in_your_estate", "frontier", "architecture")]
    assert "SECRET-SCAFFOLD-02" in pages[1]["record"] and "Slide 03" not in pages[1]["record"] and "Part 2" not in pages[1]["record"]
    assert pages[2]["title"] == "Everything runs in your tenancy" and "Speaker Notes" not in pages[2]["record"]


def test_speaker_notes_are_optional_and_never_generated_from_a_disabled_spec(tmp_path):
    runner = _runner()
    disabled = SPEC + "\n- **Generation**: disabled\n"
    enabled = SPEC + "\n- **Generation**: enabled\n"
    assert not runner.speaker_notes_enabled(SPEC)
    assert not runner.speaker_notes_enabled(disabled)
    assert runner.speaker_notes_enabled(enabled)

    (tmp_path / "design_spec.md").write_text(disabled, encoding="utf-8")
    instance = runner.Runner.__new__(runner.Runner)
    instance.project = tmp_path
    assert instance.write_notes() == 0
    assert not (tmp_path / "notes").exists()

    (tmp_path / "design_spec.md").write_text(enabled, encoding="utf-8")
    assert instance.write_notes() == 3
    assert "Message: a proposal." in (tmp_path / "notes" / "01_cover.md").read_text(encoding="utf-8")


def test_shared_system_prompt_carries_the_deck_but_no_page_record(tmp_path):
    runner = _runner()
    project = runner.ROOT / "projects" / "_test_deck_runner"
    project.mkdir(parents=True, exist_ok=True)
    try:
        (project / "design_spec.md").write_text(SPEC, encoding="utf-8")
        (project / "spec_lock.md").write_text("## canvas\n- viewBox: 0 0 1280 720\n", encoding="utf-8")
        catalog = project / "analysis" / "source_visuals" / "catalog.md"
        catalog.parent.mkdir(parents=True, exist_ok=True)
        catalog.write_text("# Source visual assets\n", encoding="utf-8")
        pages = runner.parse_pages(SPEC)
        system = runner.build_system(project, pages, "calibration table")
        assert "design system here" in system and "one loop" in system and "viewBox: 0 0 1280 720" in system and "calibration table" in system
        assert "analysis/source_visuals/catalog.md" in system
        assert "P03 (architecture, file 03_architecture_in_your_estate.svg): Everything runs in your tenancy" in system
        assert "SECRET-SCAFFOLD-02" not in system  # a page's record reaches only its own author
        task = runner.page_task(project, pages[1], pages[1])
        other = runner.page_task(project, pages[2], pages[1])
        assert "you ARE the deck's chrome anchor" in task and "SECRET-SCAFFOLD-02" in task
        assert "svg_output/02_executive_summary.svg" in other and "SECRET-SCAFFOLD-02" not in other
    finally:
        import shutil
        shutil.rmtree(project, ignore_errors=True)


def test_parallel_writers_never_lose_a_journal_entry(tmp_path):
    import threading
    project = _project(tmp_path)

    def writer(index: int) -> None:
        for step in range(15):
            page_review.update_journal(project, lambda journal, i=index: page_review._page_entry(journal, f"{i:02d}_page").__setitem__(
                "renders", page_review._page_entry(journal, f"{i:02d}_page")["renders"] + 1))

    threads = [threading.Thread(target=writer, args=(i,)) for i in range(8)]
    [t.start() for t in threads]
    [t.join() for t in threads]
    pages = page_review.load_journal(project)["pages"]
    assert len(pages) == 8 and all(entry["renders"] == 15 for entry in pages.values())
    assert not (project / ".review" / "journal.lock").exists()


def _text(index, rect, text="Label", order=10, fill="#1F1F1F", size=12):
    return {"index": index, "id": None, "order": order, "fill": fill, "size": size, "waiver": None, "rect": rect, "text": text,
            "lines": [{"rect": rect, "text": text, "fill": fill}]}


def _shape(index, rect, order=1, fill="#FFFFFF", framed=True):
    return {"index": index, "id": None, "tag": "rect", "order": order, "fill": fill, "opacity": 1.0, "framed": framed, "waiver": None, "groups": [], "rect": rect}


def _line(index, pts, order=5):
    return {"index": index, "id": None, "tag": "line", "order": order, "opacity": 1.0, "dashed": False, "waiver": None, "groups": [], "pts": pts}


def _kinds(texts=(), shapes=(), lines=(), monkeypatch=None):
    if monkeypatch is not None:
        monkeypatch.setattr(page_lint, "text_lines", lambda text: text["lines"])
    return [(f["kind"], f["hard"]) for f in page_lint.analyse({"canvas": [1280, 720], "texts": list(texts), "shapes": list(shapes), "lines": list(lines)})]


def test_lint_text_on_text_is_certain_and_a_clear_page_is_clean(monkeypatch):
    assert _kinds([_text(0, [100, 100, 300, 120]), _text(1, [100, 140, 300, 160])], monkeypatch=monkeypatch) == []
    assert ("TEXT_ON_TEXT", True) in _kinds([_text(0, [100, 100, 300, 120], "Pilot window"), _text(1, [180, 104, 380, 124], "Week 12")], monkeypatch=monkeypatch)


def test_lint_line_through_text_is_flagged_but_a_tether_is_not(monkeypatch):
    label = _text(0, [200, 200, 320, 218])
    through = _line(0, [[x, 209] for x in range(100, 420, 5)])
    tether = _line(1, [[x, 209] for x in range(100, 206, 5)])  # ends at the label's edge
    assert ("LINE_THROUGH_TEXT", False) in _kinds([label], lines=[through], monkeypatch=monkeypatch)
    assert _kinds([label], lines=[tether], monkeypatch=monkeypatch) == []


def test_lint_a_grid_line_behind_a_bar_does_not_touch_the_label_on_it(monkeypatch):
    bar = _shape(0, [150, 195, 400, 225], order=6, fill="#B4162E", framed=False)
    label = _text(0, [200, 200, 320, 218], order=10, fill="#FFFFFF")
    grid = _line(0, [[260, y] for y in range(100, 400, 5)], order=2)
    assert _kinds([label], shapes=[bar], lines=[grid], monkeypatch=monkeypatch) == []


def test_lint_text_overflowing_its_box_is_flagged(monkeypatch):
    box = _shape(0, [100, 100, 260, 140])
    assert ("TEXT_OFF_SHAPE", False) in _kinds([_text(0, [110, 110, 300, 128])], shapes=[box], monkeypatch=monkeypatch)
    assert _kinds([_text(0, [110, 110, 250, 128])], shapes=[box], monkeypatch=monkeypatch) == []


def test_lint_off_canvas_is_certain(monkeypatch):
    assert ("OFF_CANVAS", True) in _kinds([_text(0, [1200, 100, 1300, 118])], monkeypatch=monkeypatch)


def _deck_with_floating_text(path):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Emu, Pt
    px = 9525
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(1280 * px), Emu(720 * px)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(100 * px), Emu(100 * px), Emu(300 * px), Emu(120 * px))
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(0xF1, 0xF2, 0xF4)

    def label(x, y, text, size=12):
        box = slide.shapes.add_textbox(Emu(x * px), Emu(y * px), Emu(160 * px), Emu(20 * px))
        box.text_frame.word_wrap = False
        for side in ("margin_left", "margin_top", "margin_right", "margin_bottom"):
            setattr(box.text_frame, side, 0)
        run = box.text_frame.paragraphs[0].add_run()
        run.text, run.font.size = text, Pt(size)

    label(116, 112, "Heading in the card")
    label(116, 140, "Body line in the card", 9)
    label(600, 300, "Free heading")            # no shape under these two: they become one text box
    label(600, 322, "Free line under it", 9)
    label(900, 600, "A lone label")
    prs.save(path)


def test_text_moves_into_the_shape_it_sits_on(tmp_path):
    import pptx_text_in_shapes
    from pptx import Presentation
    source, target = tmp_path / "in.pptx", tmp_path / "out.pptx"
    _deck_with_floating_text(source)
    totals = pptx_text_in_shapes.convert(source, target)["totals"]
    assert (totals["texts"], totals["adopted"], totals["hosts"], totals["stacked"], totals["stacks"], totals["single"]) == (5, 2, 1, 2, 1, 1)
    shapes = list(Presentation(target).slides[0].shapes)
    card = next(s for s in shapes if s.shape_type != 17)
    assert [p.text for p in card.text_frame.paragraphs] == ["Heading in the card", "Body line in the card"]
    assert card.text_frame.word_wrap is True and card.text_frame.margin_left == 16 * 9525 and card.text_frame.margin_top == 12 * 9525
    boxes = sorted((s for s in shapes if s.shape_type == 17), key=lambda s: s.left)
    assert [[p.text for p in b.text_frame.paragraphs] for b in boxes] == [["Free heading", "Free line under it"], ["A lone label"]]


def _runner_module():
    import importlib
    host_dir = Path(__file__).resolve().parents[4] / "hosts" / "responses_api"
    if str(host_dir) not in sys.path:
        sys.path.insert(0, str(host_dir))
    return importlib.import_module("deck_runner")


def test_a_repair_never_demotes_a_page_that_had_passed():
    runner = _runner_module()
    passed = {"outcome": "accepted", "last_sha": "a", "review": {"sha": "a", "verdict": "PASS", "blockers": 0}, "lint": {"hard": 0}}
    polish = {"outcome": "unresolved", "last_sha": "b", "review": {"sha": "b", "verdict": "EXECUTION_REPAIR", "blockers": 0}, "lint": {"hard": 0}}
    broken = {"outcome": "unresolved", "last_sha": "c", "review": {"sha": "c", "verdict": "EXECUTION_REPAIR", "blockers": 2}, "lint": {"hard": 0}}
    collided = {"outcome": "unresolved", "last_sha": "d", "review": {"sha": "d", "verdict": "EXECUTION_REPAIR", "blockers": 0}, "lint": {"hard": 1}}
    unreviewed = {"outcome": None, "last_sha": "e", "review": {"sha": "a", "verdict": "PASS", "blockers": 0}, "lint": {"hard": 0}}
    replan = {"outcome": "unresolved", "last_sha": "f", "review": {"sha": "f", "verdict": "CONCEPT_REPLAN", "blockers": 0}, "lint": {"hard": 0}}
    assert runner.keep_or_restore(passed, {**passed, "last_sha": "g", "review": {"sha": "g", "verdict": "PASS", "blockers": 0}}) == "fine"
    assert runner.keep_or_restore(passed, polish) == "keep"          # no defect, only polish: the repaired revision stays accepted
    assert runner.keep_or_restore(passed, broken) == "restore"       # real defects: the revision that passed comes back
    assert runner.keep_or_restore(passed, collided) == "restore"
    assert runner.keep_or_restore(passed, unreviewed) == "restore"   # the repair was never reviewed
    assert runner.keep_or_restore(passed, replan) == "restore"
    assert runner.keep_or_restore({"outcome": "unresolved"}, broken) == "fine"  # it had not passed: nothing to protect


def test_comments_from_the_preview_are_collected_and_addressed_to_their_page(tmp_path):
    runner = _runner_module()
    (tmp_path / "svg_output").mkdir()
    (tmp_path / "svg_output" / "03_plan.svg").write_text(
        '<svg xmlns="http://www.w3.org/2000/svg"><g id="lane-2"><text id="bar-label-4" data-edit-target="true" '
        'data-edit-annotation="Shorten to two words and put it inside the bar">PR mode on pilot services</text></g>'
        '<rect id="gate" data-edit-target="true" data-edit-annotation="  "/></svg>', encoding="utf-8")
    (tmp_path / "svg_output" / "04_team.svg").write_text('<svg xmlns="http://www.w3.org/2000/svg"><text id="t">clean</text></svg>', encoding="utf-8")
    found = runner.collect_annotations(tmp_path)
    assert list(found) == ["03_plan"] and found["03_plan"] == [{"element": "bar-label-4", "tag": "text", "text": "PR mode on pilot services",
                                                                 "comment": "Shorten to two words and put it inside the bar"}]
    message = runner.revision_message("03_plan", found["03_plan"])
    assert 'id="bar-label-4"' in message and "Shorten to two words" in message and "data-edit-annotation" in message and "DONE 03_plan" in message


def _bare_runner(runner, **args):
    import argparse
    bare = object.__new__(runner.Runner)
    bare.args = argparse.Namespace(repair_rounds=2, no_escalate=False, max_parallel=4, **args)
    bare.page_sessions, bare.authored_by, bare.checker_repairs, bare.session_tier = {}, {}, {}, {}
    bare.say = lambda text: None
    return bare


def test_a_page_is_repaired_as_soon_as_its_author_ends_not_after_its_slowest_sibling():
    import threading
    import time as _time
    runner = _runner_module()
    bare = _bare_runner(runner)
    events, lock = [], threading.Lock()
    pending = {"fast": [["a blocking item"]], "slow": []}

    def author(page, tier, anchor, **_):
        _time.sleep(0.6 if page["stem"] == "slow" else 0.05)
        bare.page_sessions.setdefault(page["stem"], []).append(page["stem"])
        bare.authored_by[page["stem"]] = tier
        with lock:
            events.append(("authored", page["stem"]))

    bare.author = author
    bare.page_checker = lambda page: pending[page["stem"]].pop(0) if pending[page["stem"]] else []
    bare.repair = lambda page, issues, hint="": events.append(("repaired", page["stem"]))
    bare.journal_page = lambda stem: {"outcome": "accepted"}
    pages = [{"stem": "fast", "number": 1, "tier": "frontier"}, {"stem": "slow", "number": 2, "tier": "frontier"}]
    bare.fan_out([lambda p=p: (bare.author(p, p["tier"], None), bare.after_author(p, None)) for p in pages])
    assert events.index(("repaired", "fast")) < events.index(("authored", "slow"))
    assert bare.checker_repairs == {"fast": 1}


def test_page_checker_repairs_share_one_budget_and_an_unaccepted_workhorse_page_escalates_in_its_own_job():
    runner = _runner_module()
    bare = _bare_runner(runner)
    page = {"stem": "p", "number": 3, "tier": "workhorse"}
    bare.page_sessions["p"], bare.authored_by["p"] = ["s"], "workhorse"
    calls = []
    bare.page_checker = lambda page: ["still broken"]
    bare.repair = lambda page, issues, hint="": calls.append("repair")
    bare.journal_page = lambda stem: {"outcome": "unresolved"}

    def escalate(page, anchor):
        calls.append("escalate")
        bare.authored_by["p"] = "frontier"

    bare.escalate = escalate
    bare.after_author(page, None)
    assert calls == ["repair", "repair", "escalate"]  # two checker repairs (the budget), escalation, and no third repair
    assert bare.checker_repairs["p"] == 2


def test_lint_stops_its_background_contract_check_when_the_measurement_fails(tmp_path, monkeypatch):
    killed = []

    class Proc:
        def kill(self):
            killed.append(True)

    scratch = tmp_path / "scratch"
    scratch.mkdir()
    (tmp_path / "svg_output").mkdir()
    (tmp_path / "svg_output" / "01_x.svg").write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720"/>', encoding="utf-8")
    monkeypatch.setattr(page_lint, "start_contract", lambda svg: {"svg": svg, "proc": Proc(), "scratch": str(scratch), "report": scratch / "r.json"})

    def broken(files, browser=None):
        raise RuntimeError("no browser")
        yield  # pragma: no cover

    monkeypatch.setattr(page_lint, "measure", broken)
    try:
        page_lint.lint_page(tmp_path, "01_x")
    except RuntimeError:
        pass
    assert killed and not scratch.exists()


def test_the_template_stage_is_not_told_the_page_authors_never_touch_templates():
    runner = _runner_module()
    project = runner.ROOT / "projects" / "_test_template_system"
    text = runner.template_brief_system(project)
    assert "Never edit the design spec, the lock, another page, the templates" not in text  # the page rule a strict model obeys
    assert "projects/_test_template_system/templates/" in text and text.startswith("# Template designer")
    assert "## The record" in text  # the rest of the page brief stays, as reference


def _record(content_words: int, fill: str = "zone A x=58..600, y=180..420; zone B x=620..1176, y=180..420; strip x=58..1176, y=430..662",
            scaffold_words: int = 120, extra_content: str = "") -> str:
    fields = {"Role": "diagram", "Author tier": "frontier", "Layout": "content", "Audience question": "How does the flow work?",
              "Audience move": "Unfamiliar to informed", "Story link": "Follows the context and prepares the example", "Relationships": "input precedes check precedes output",
              "Title": "A finding", "Core message": "Why it matters", "Content": " ".join(["word"] * content_words) + extra_content,
              "Sources": "sources/request.md", "Visual task": "Trace input through check to output", "Visual approach": "A labelled process because order matters",
              "Composition": "Input at left, check as focal center, output at right; read left to right",
              "Hierarchy": "Flow first, explanation second", "Avoid": "unlabelled arrows", "Editor notes": "none"}
    return "#### Slide 02 - Example\n\n" + "\n".join(f"- **{k}**: {v}" for k, v in fields.items()) + "\n"


def test_the_plan_check_accepts_a_spare_record_and_requires_a_reader_question_and_visual_task():
    runner = _runner_module()
    assert runner.plan_lint("# spec\n\n" + _record(18), 330) == {}
    missing = _record(18).replace("- **Audience question**: How does the flow work?", "- **Audience question**: ")
    missing = missing.replace("- **Visual task**: Trace input through check to output", "- **Visual task**: ")
    issues = " | ".join(runner.plan_lint("# spec\n\n" + missing)["Slide 02 - Example"])
    assert "Audience question" in issues and "Visual task" in issues


def test_the_plan_check_rejects_pixel_geometry_in_semantic_visual_guidance():
    runner = _runner_module()
    record = _record(18).replace("A labelled process because order matters", "Three boxes x=50 y=200 because order matters")
    issues = " | ".join(runner.plan_lint("# spec\n\n" + record)["Slide 02 - Example"])
    assert "element geometry" in issues


def test_slide_roster_accepts_standard_and_typographic_heading_dashes():
    runner = _runner_module()
    spec = "## IX. Content Outline\n#### Slide 01 - First\n- **Title**: One\n#### Slide 02 — Second\n- **Title**: Two\n"
    pages = runner.parse_pages(spec)
    assert [(page["number"], page["name"], page["title"]) for page in pages] == [(1, "First", "One"), (2, "Second", "Two")]


def test_the_plan_check_catches_filler_repeated_rows_and_internal_notes():
    runner = _runner_module()
    filler = "\n    " + "\n    ".join(["cell: `To assess`"] * 6)
    rows = "\n    " + "\n    ".join(["row: named owner, scope, phased timing and dependency"] * 3)
    leak = "\n    note: `Submission blocker - deck alone is not ready to file`"
    issues = " | ".join(runner.plan_lint("# spec\n\n" + _record(400, extra_content=filler + rows + leak), 330)["Slide 02 - Example"])
    assert "filler" in issues and "identical rows" in issues and "internal note" in issues


def test_the_exporters_table_gate_runs_while_the_page_is_drawn(tmp_path):
    class Done:
        def __init__(self, out):
            self.out = out

        def communicate(self, timeout=None):
            return self.out, ""

        def kill(self):
            pass

    found = page_lint._native_findings(Done('[{"marker": "fees-table", "finding": "fallback text is missing from metadata"}, {"marker": "fees-table", "finding": "colour not projected"}]'))
    assert [f["kind"] for f in found] == ["NATIVE", "NATIVE"] and all(f["hard"] for f in found)
    assert "fees-table: fallback text" in found[0]["message"] and "stamp_native_fallbacks.py" in found[0]["message"]  # the fix travels with the first item
    assert "stamp_native_fallbacks.py" not in found[1]["message"]
    assert page_lint._native_findings(None) == [] and page_lint._native_findings(Done("not json")) == []
    plain = tmp_path / "01_x.svg"
    plain.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720"><text>no table</text></svg>', encoding="utf-8")
    import subprocess
    out = subprocess.run([sys.executable, str(Path(page_lint.__file__).with_name("native_parity.py")), str(plain)], capture_output=True, text=True)
    assert out.stdout.strip() == "[]"

def test_the_plan_check_names_what_a_reader_with_no_context_cannot_read():
    runner = _runner_module()
    spec = ("# spec\n\n#### Slide 02 - Why now\n\n- **Title**: The RFP asks for a ROM estimate\n- **Core message**: Everything by W23\n"
            "- **Content**:\n    eyebrow: `WHY THIS MATTERS NOW`\n    body: `Our Trust Gate (TG) checks data before release; see §3.2 of the workbook`\n"
            "- **Fill**: x\n\n#### Slide 03 - Later\n\n- **Title**: TG blocks bad data\n- **Content**:\n    body: `Request for proposal (RFP) answered; week 23 (W23) of the engagement`\n")
    found = runner.undefined_terms(spec)
    first = " | ".join(found["Slide 02 - Why now"])
    assert "RFP" in first and "ROM" in first and "TG" not in first  # TG is defined on the page, "Trust Gate (TG)"
    assert "WHY" not in first and "MATTERS" not in first  # an eyebrow set in capitals is typography, not an abbreviation
    assert "week codes" in first and "reference to material" in first
    assert "Slide 03 - Later" not in found  # RFP is reported once, at its first use; W23 is explained as a week there

def test_a_template_s_own_fonts_win_over_a_default_office_theme(tmp_path):
    import zipfile
    runner = _runner_module()
    default_theme = ('<a:theme><a:majorFont><a:latin typeface="Calibri"/></a:majorFont><a:minorFont><a:latin typeface="Calibri"/></a:minorFont>'
                     '<a:accent1><a:srgbClr val="4F81BD"/></a:accent1></a:theme>')
    slide = ('<p:sld><p:sp><p:nvSpPr><p:nvPr><p:ph type="title"/></p:nvPr></p:nvSpPr><a:rPr><a:latin typeface="Georgia"/></a:rPr><a:srgbClr val="182B59"/></p:sp>'
             '<p:sp><a:rPr><a:latin typeface="Arial"/></a:rPr></p:sp></p:sld>')
    deck = tmp_path / "t.pptx"
    with zipfile.ZipFile(deck, "w") as z:
        z.writestr("ppt/theme/theme1.xml", default_theme)
        z.writestr("ppt/slides/slide1.xml", slide)
    used = runner.template_styles_used(deck)
    assert used["title_fonts"] == ["Georgia"] and used["text_fonts"] == ["Arial"] and "#182B59" in used["colours"] and used["theme_is_default"]
    styled = tmp_path / "s.pptx"
    with zipfile.ZipFile(styled, "w") as z:
        z.writestr("ppt/theme/theme1.xml", default_theme.replace("Calibri", "Raleway"))
        z.writestr("ppt/theme/theme2.xml", default_theme)  # a notes theme left at the Office default
        z.writestr("ppt/slides/slide1.xml", slide)
    assert runner.template_styles_used(styled)["theme_is_default"] is False

def test_a_wrong_page_number_in_the_footer_is_a_certain_defect():
    def geometry(value):
        return {"canvas": [1280, 720], "texts": [{"text": value, "parts": [{"rect": [1148, 664, 1160, 684]}]},
                                                 {"text": "12", "parts": [{"rect": [300, 400, 320, 420]}]}]}  # a number in the body is not the folio
    wrong = page_lint.folio_issues(geometry("7"), "08_what_the_checks_see")
    assert len(wrong) == 1 and wrong[0]["kind"] == "FOLIO" and wrong[0]["hard"] and "reads 7" in wrong[0]["message"]
    assert page_lint.folio_issues(geometry("8"), "08_what_the_checks_see") == []
    assert page_lint.folio_issues(geometry("8"), "cover") == []  # no number in the stem: nothing to compare


def test_consulting_source_intake_exposes_embedded_and_supplied_images(tmp_path, monkeypatch):
    import subprocess
    from PIL import Image
    from pptx import Presentation
    sys.path.insert(0, str(Path(__file__).resolve().parents[4] / "hosts" / "responses_api"))
    import source_assets

    sources = tmp_path / "sources"
    sources.mkdir()
    photo = sources / "sample.png"
    Image.new("RGB", (64, 48), "#146F77").save(photo)
    deck = sources / "example.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(str(photo), 0, 0)
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(deck)

    original_run = source_assets._run

    def render_fixture(script, *args):
        if script.name != "pptx_render.py":
            return original_run(script, *args)
        output = Path(args[args.index("--out") + 1])
        output.mkdir(parents=True, exist_ok=True)
        for number in range(1, len(Presentation(args[0]).slides) + 1):
            Image.new("RGB", (160, 90), "white").save(output / f"slide-{number:03d}.png")
        return subprocess.CompletedProcess([str(script)], 0, "rendered", "")

    monkeypatch.setattr(source_assets, "_run", render_fixture)
    scripts = Path(__file__).resolve().parents[1]
    first = source_assets.prepare(tmp_path, scripts)
    second = source_assets.prepare(tmp_path, scripts)
    kinds = {item["kind"] for item in first["assets"]}
    assert kinds == {"embedded_image", "supplied_image", "source_slide_preview"}
    assert len(second["assets"]) == len(first["assets"]) == 4
    assert any(item["source"] == "sources/example.pptx" and item["kind"] == "embedded_image" for item in first["assets"])
    assert all((tmp_path / item["path"]).is_file() for item in first["assets"])
    assert (tmp_path / "analysis" / "source_visuals" / "catalog.md").is_file()
    catalog = (tmp_path / "analysis" / "source_visuals" / "catalog.md").read_text(encoding="utf-8")
    assert "on source slide 1" in catalog and "slide 2" in catalog

    revised = Presentation()
    revised.slides.add_slide(revised.slide_layouts[6])
    revised.save(deck)
    latest = source_assets.prepare(tmp_path, scripts)
    assert {item["kind"] for item in latest["assets"]} == {"supplied_image", "source_slide_preview"}
    assert [item["slide"] for item in latest["assets"] if item["kind"] == "source_slide_preview"] == [1]
    assert "slide 2" not in (tmp_path / "analysis" / "source_visuals" / "catalog.md").read_text(encoding="utf-8")
