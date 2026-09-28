"""Diagram and timeline tools (failure families F04 and F05): the diagram lint rules, the diagram contract's delivery to authors and
reviewers, the ELK diagram layout helper, the Gantt layout helper and the elbow connectors made at export. No model is called."""

import json
import os
import shutil
import sys
from pathlib import Path

import pytest

os.environ["PPT_MASTER_NO_LINT"] = "1"
SCRIPTS = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCRIPTS))

import page_lint  # noqa: E402
import timeline_layout  # noqa: E402


# ---------------------------------------------------------------------------------------------------------------- lint rules

def _text(index, rect, text="Label", order=50, size=12, ph=None):
    return {"index": index, "id": None, "order": order, "fill": "rgb(20, 20, 20)", "size": size, "waiver": None, "rect": rect, "text": text,
            "ph": ph, "lines": [{"rect": rect, "text": text, "fill": "rgb(20, 20, 20)"}]}


def _box(index, rect, order=1, framed=True, fill="rgb(255, 255, 255)", tag="rect", npts=None):
    return {"index": index, "id": None, "tag": tag, "order": order, "fill": fill, "opacity": 1.0, "framed": framed, "waiver": None,
            "groups": [], "rect": rect, "npts": npts, "stroke": "rgb(0, 0, 0)" if framed else None}


def _line(index, pts, order=5, heads=(False, False), dashed=False, stroke="rgb(191, 63, 26)"):
    return {"index": index, "id": None, "tag": "line", "order": order, "opacity": 1.0, "dashed": dashed, "waiver": None, "groups": [],
            "pts": [list(p) for p in pts], "heads": list(heads), "stroke": stroke}


def _head(index, tip, pointing="up"):
    x, y = tip
    rect = {"up": [x - 4, y, x + 4, y + 8], "down": [x - 4, y - 8, x + 4, y], "right": [x - 8, y - 4, x, y + 4]}[pointing]
    return _box(index, rect, order=40, framed=False, fill="rgb(191, 63, 26)", tag="polygon", npts=3)


def _findings(texts=(), shapes=(), lines=(), monkeypatch=None):
    monkeypatch.setattr(page_lint, "text_lines", lambda text: text["lines"])
    geometry = {"canvas": [1280, 720], "texts": list(texts), "shapes": list(shapes), "lines": list(lines)}
    return [f["kind"] for f in page_lint.analyse(geometry)]


def _row_of_steps():
    """Six steps in a row with their labels inside, like S6 P09 (1280 x 720 page)."""
    boxes, labels = [], []
    for i in range(6):
        x = 82 + i * 186
        boxes.append(_box(i, [x, 275, x + 172, 391]))
        labels.append(_text(i, [x + 10, 318, x + 110, 334], f"Step {i + 1}"))
    return boxes, labels


def test_an_orphan_label_in_the_figure_is_flagged_and_an_attached_one_is_not(monkeypatch):
    boxes, labels = _row_of_steps()
    loop = [_line(90, [[1144, 391], [1144, 456]]), _line(91, [[1144, 456], [168, 456]]), _line(92, [[168, 456], [168, 391]])]
    head = _head(80, (168, 391), "up")
    orphan = _text(20, [741, 483, 951, 491], "PASS -> ACCEPTED BUDGET -> UNRESOLVED", size=11)
    on_line = _text(21, [480, 441, 700, 452], "CONCEPT_REPLAN -> PLAN + WRITE", size=11)  # 4 px above its route
    kinds = _findings(labels + [orphan, on_line], boxes + [head], loop, monkeypatch)
    assert kinds.count("ORPHAN_LABEL") == 1
    assert "LOOP_WITHOUT_HEAD" not in kinds and "CONNECTOR_NO_TARGET" not in kinds


def test_a_return_route_without_an_arrowhead_is_flagged_and_with_one_it_is_not(monkeypatch):
    boxes, labels = _row_of_steps()
    loop = [_line(90, [[1077, 391], [1077, 423]]), _line(91, [[1077, 423], [355, 423]]), _line(92, [[355, 423], [355, 391]])]
    assert "LOOP_WITHOUT_HEAD" in _findings(labels, boxes, loop, monkeypatch)
    assert "LOOP_WITHOUT_HEAD" not in _findings(labels, boxes + [_head(80, (355, 391), "up")], loop, monkeypatch)
    marker = [_line(90, [[1077, 391], [1077, 423]]), _line(91, [[1077, 423], [355, 423]]), _line(92, [[355, 423], [355, 391]], heads=(False, True))]
    assert "LOOP_WITHOUT_HEAD" not in _findings(labels, boxes, marker, monkeypatch)


def test_a_labelled_underline_under_a_row_of_steps_stands_in_for_a_loop(monkeypatch):
    boxes, labels = _row_of_steps()  # S6 first export P09: the loops were drawn as captioned underlines
    underline = _line(90, [[455, 425], [1188, 425]])
    caption = _text(20, [461, 409, 700, 421], "EXECUTION_REPAIR -> fix and rerender", size=11)
    assert "LOOP_WITHOUT_HEAD" in _findings(labels + [caption], boxes, [underline], monkeypatch)
    grey_rule = _line(90, [[455, 425], [1188, 425]], stroke="rgb(216, 213, 206)")  # a pale separator is a rule, not a route
    assert "LOOP_WITHOUT_HEAD" not in _findings(labels + [caption], boxes, [grey_rule], monkeypatch)


def test_an_arrow_that_points_at_empty_space_is_flagged(monkeypatch):
    boxes, labels = _row_of_steps()
    to_nothing = _line(90, [[1064, 391], [1064, 440]], heads=(False, True))
    into_box = _line(91, [[254, 331], [268, 331]], heads=(False, True))
    kinds = _findings(labels, boxes, [to_nothing, into_box], monkeypatch)
    assert kinds.count("CONNECTOR_NO_TARGET") == 1


def test_numbered_notes_need_matching_markers_on_the_figure(monkeypatch):
    boxes = [_box(i, [600, 150 + i * 110, 900, 230 + i * 110]) for i in range(3)]
    labels = [_text(i, [620, 180 + i * 110, 760, 196 + i * 110], f"Component {i}") for i in range(3)]
    notes = _text(30, [72, 320, 440, 440], "notes", size=15, ph="body")
    notes["lines"] = [{"rect": [72, 320 + k * 40, 440, 336 + k * 40], "text": f"{k + 1} The note about component {k}", "fill": "x"} for k in range(3)]
    kinds = _findings(labels + [notes], boxes, [_line(90, [[750, 230], [750, 260]], heads=(False, True))], monkeypatch)
    assert "UNKEYED_CALLOUT" in kinds
    badges = [_box(10 + k, [590, 140 + k * 110, 612, 162 + k * 110], order=60, framed=False, fill="rgb(191, 63, 26)", tag="circle") for k in range(3)]
    numerals = [_text(40 + k, [596, 145 + k * 110, 606, 158 + k * 110], str(k + 1), order=61, size=11) for k in range(3)]
    kinds = _findings(labels + [notes] + numerals, boxes + badges, [_line(90, [[750, 230], [750, 260]], heads=(False, True))], monkeypatch)
    assert "UNKEYED_CALLOUT" not in kinds


def test_a_label_cut_by_a_container_outline_is_flagged(monkeypatch):
    machine = _box(0, [514, 153, 1011, 573], fill=None)  # S6 first export P05: "prompts / text / page" across YOUR MACHINE
    inside = _box(1, [711, 203, 984, 266])
    kinds = _findings([_text(0, [1002, 182, 1129, 195], "prompts / text / page", size=11), _text(1, [721, 220, 820, 236], "Agent host")],
                      [machine, inside], [], monkeypatch)
    assert kinds.count("BOUNDARY_CROSSING") == 1


def test_diagram_rules_leave_a_plain_text_page_alone(monkeypatch):
    texts = [_text(i, [72, 150 + 30 * i, 600, 166 + 30 * i], f"Paragraph line {i} of a plain page") for i in range(8)]
    assert _findings(texts, [], [], monkeypatch) == []


# ---------------------------------------------------------------------------------------------------- contract delivery

def _runner_module():
    import importlib
    host_dir = Path(__file__).resolve().parents[4] / "hosts" / "responses_api"
    if str(host_dir) not in sys.path:
        sys.path.insert(0, str(host_dir))
    return importlib.import_module("deck_runner")


def test_authors_and_reviewers_receive_the_diagram_contract():
    import page_review
    runner = _runner_module()
    assert "references/diagram-clarity.md" in runner.CONTRACT_DOCS
    assert "diagram_layout.py" in runner.PAGE_AUTHOR_BRIEF and "timeline_layout.py" in runner.PAGE_AUTHOR_BRIEF
    contract = page_review._diagram_contract()
    assert contract.startswith("## Diagram contract")
    for phrase in ("arrowhead", "loop", "orphan", "callout", "axis", "boundary"):
        assert phrase in contract.lower()
    assert len(contract.splitlines()) <= 60
    assert "DIAGRAM CONTRACT" in page_review.REVIEW_INSTRUCTIONS


# ------------------------------------------------------------------------------------------------------- timeline layout

def _gantt(lanes: int, weeks: int) -> dict:
    words = ["Discovery", "Interviews and workshops", "Baseline", "Design", "Build and test", "Pilot", "Rollout", "Training", "Handover",
             "Security review", "Data contracts", "Reporting"]
    spec_lanes = []
    for lane in range(lanes):
        bars = []
        start = 1 + (lane * 3) % max(1, weeks // 3)
        for k in range(3):
            length = 2 + (lane + k) % 6
            end = min(weeks, start + length)
            bars.append({"id": f"b{lane}_{k}", "start": start, "end": end, "label": words[(lane + k) % len(words)]})
            start = end + 1 + (k % 2) * 3
            if start > weeks - 2:
                break
        spec_lanes.append({"id": str(lane), "label": f"Workstream {lane + 1}", "bars": bars})
    step = max(2, weeks // 6)
    return {"region": {"x": 50, "y": 110, "w": 1180, "h": 580}, "horizon": {"unit": "week", "start": 1, "end": weeks, "prefix": "W"},
            "lane_label_w": 200, "lanes": spec_lanes,
            "gates": [{"at": w, "label": f"Gate {i + 1}"} for i, w in enumerate(range(step, weeks, step * 2))],
            "milestones": [{"at": w, "label": f"Milestone {i + 1}"} for i, w in enumerate(range(2, weeks, step))],
            "dependencies": [{"from": "b0_0", "to": "b1_1"}, {"from": "b2_0", "to": "b3_1"}]}


def test_bars_and_markers_sit_exactly_on_the_week_scale():
    placed = timeline_layout.layout(_gantt(9, 26))
    scale = placed["scale"]
    unit_w = (1230 - scale["x0"]) / 26
    assert abs(scale["unit_w"] - unit_w) < 1e-3
    for bar in placed["bars"]:
        assert abs(bar["x0"] - (scale["x0"] + (bar["start"] - 1) * unit_w)) < 0.02
        assert abs(bar["x1"] - (scale["x0"] + bar["end"] * unit_w)) < 0.02  # a bar ends at the right edge of its last week
    for marker in placed["gates"] + placed["milestones"]:
        assert abs(marker["x"] - (scale["x0"] + marker["at"] * unit_w)) < 0.02  # "end" convention: at the end of week N
    assert timeline_layout.Scale(1, 10, 100, 500, "middle").marker(3) == 100 + 2.5 * 50


@pytest.mark.parametrize("lanes,weeks", [(9, 26), (6, 51)])
def test_no_two_labels_collide_and_every_bar_is_named(lanes, weeks):
    placed = timeline_layout.layout(_gantt(lanes, weeks))
    rects = timeline_layout.label_rects(placed)
    for i, (name_a, a) in enumerate(rects):
        for name_b, b in rects[i + 1:]:
            assert not timeline_layout._overlap(a, b, -0.5), f"{name_a} overlaps {name_b}"
    assert all(b["text"]["lines"] and " ".join(b["text"]["lines"]).replace(" ", "") == b["label"].replace(" ", "") for b in placed["bars"])  # never truncated
    assert not [c for c in placed["checks"] if "overlaps" in c or "straddles" in c or "runs out" in c]
    assert min(placed["font"][k] for k in ("label_px", "lane_px", "tick_px")) >= 14


def test_a_label_goes_inside_when_it_fits_else_beside_else_above():
    scale = (100.0, 1200.0)
    wide = timeline_layout.place_bar_label({"x0": 200, "x1": 500, "label": "Build"}, 40, 14, scale, "Segoe UI", 300)
    beside = timeline_layout.place_bar_label({"x0": 200, "x1": 230, "label": "Interviews"}, 70, 14, scale, "Segoe UI", 30)
    before = timeline_layout.place_bar_label({"x0": 1150, "x1": 1190, "label": "Handover"}, 60, 14, scale, "Segoe UI", 40)
    arrow_out = timeline_layout.place_bar_label({"x0": 1150, "x1": 1190, "label": "Handover"}, 60, 14, scale, "Segoe UI", 40, forbid=("left",))
    assert (wide["where"], beside["where"], before["where"], arrow_out["where"]) == ("inside", "right", "left", "above")
    assert beside["x"] == 236


def test_tick_labels_are_thinned_until_none_collide():
    scale = timeline_layout.Scale(1, 51, 250, 980)
    step, ticks = timeline_layout.thin_ticks(scale, "week", "W", 14, "Segoe UI")
    assert step > 1 and all(b["rect"][0] - a["rect"][2] >= 8 for a, b in zip(ticks, ticks[1:]))


# -------------------------------------------------------------------------------------------------------- diagram layout

needs_node = pytest.mark.skipif(not (os.environ.get("PPT_MASTER_NODE") or shutil.which("node")), reason="Node.js is not installed")


@needs_node
def test_a_loop_is_laid_out_orthogonally_with_its_arrow_on_the_node_it_returns_to():
    import diagram_layout
    spec = {"region": {"x": 60, "y": 150, "w": 1160, "h": 420}, "direction": "RIGHT",
            "nodes": [{"id": "a", "label": "Write"}, {"id": "b", "label": "Render"}, {"id": "c", "label": "Review"},
                      {"id": "ok", "label": "Accepted", "kind": "end"}],
            "edges": [{"source": "a", "target": "b"}, {"source": "b", "target": "c"}, {"source": "c", "target": "ok", "label": "pass"},
                      {"source": "c", "target": "b", "label": "repair", "kind": "return"}]}
    placed = diagram_layout.layout(spec)
    assert placed["fits"] and placed["checks"] == [] and placed["font"]["node_px"] >= 16 and placed["font"]["label_px"] >= 14
    nodes = placed["nodes"]
    for edge in placed["edges"]:
        pts = edge["points"]
        assert all(abs(p[0] - q[0]) < 0.2 or abs(p[1] - q[1]) < 0.2 for p, q in zip(pts, pts[1:]))  # orthogonal
        target = nodes[edge["target"]]
        end = pts[-1]
        on_edge = min(abs(end[0] - target["x"]), abs(end[0] - target["x"] - target["w"]), abs(end[1] - target["y"]), abs(end[1] - target["y"] - target["h"]))
        assert on_edge < 1.0  # the arrowhead lands on the target's edge
    loop = next(e for e in placed["edges"] if e["kind"] == "return")
    assert abs(loop["points"][-1][1] - (nodes["b"]["y"] + nodes["b"]["h"])) < 1.0  # it re-enters from below, under the row
    defs, group = diagram_layout.svg_fragment(placed, spec)
    assert 'marker-end="url(#dg-arrow-accent)"' in group and "<path" in group and "<marker" in defs


@needs_node
def test_a_diagram_that_cannot_fit_at_the_type_floors_says_so_instead_of_shrinking():
    import diagram_layout
    spec = {"region": {"x": 60, "y": 150, "w": 300, "h": 120},
            "nodes": [{"id": f"n{i}", "label": f"Stage number {i}"} for i in range(8)],
            "edges": [{"source": f"n{i}", "target": f"n{i + 1}"} for i in range(7)]}
    placed = diagram_layout.layout(spec)
    assert not placed["fits"] and placed["font"]["node_px"] >= 16 and any("does not fit" in c for c in placed["checks"])


def test_a_missing_node_runtime_is_a_clear_error(monkeypatch):
    import diagram_layout
    monkeypatch.delenv("PPT_MASTER_NODE", raising=False)
    monkeypatch.setattr(diagram_layout.shutil, "which", lambda name: None)
    with pytest.raises(diagram_layout.LayoutError, match="Node.js"):
        diagram_layout.node_executable()


# ---------------------------------------------------------------------------------------------------- elbow connectors

def test_elbow_routes_map_onto_connector_presets_exactly():
    import pptx_text_in_shapes as t
    px = t.PX
    z = [(100 * px, 100 * px), (160 * px, 100 * px), (160 * px, 300 * px), (400 * px, 300 * px)]  # H V H, left to right, downwards
    prst, off, ext, rot, fh, fv, adj = t.elbow_connector(z)
    assert (prst, rot, fh, fv) == ("bentConnector3", 0, False, False) and adj == 20000 and off == (100 * px, 100 * px)
    vhv = [(400 * px, 100 * px), (400 * px, 180 * px), (100 * px, 180 * px), (100 * px, 300 * px)]  # V H V, right to left
    found = t.elbow_connector(vhv)
    assert found and found[0] == "bentConnector3" and found[3] == 5400000
    ell = [(100 * px, 100 * px), (300 * px, 100 * px), (300 * px, 50 * px)]
    assert t.elbow_connector(ell)[0] == "bentConnector2"
    u_turn = [(500 * px, 391 * px), (500 * px, 423 * px), (200 * px, 423 * px), (200 * px, 391 * px)]  # level ends: stays a drawing
    assert t.elbow_connector(u_turn) is None
    diagonal = [(0, 0), (100 * px, 50 * px), (200 * px, 50 * px)]
    assert t.elbow_connector(diagonal) is None


def _deck_with_elbow(path):
    from lxml import etree
    from pptx import Presentation
    from pptx.util import Emu
    px = 9525
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(1280 * px), Emu(720 * px)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    from pptx.enum.shapes import MSO_SHAPE
    a = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(100 * px), Emu(100 * px), Emu(120 * px), Emu(60 * px))
    b = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(400 * px), Emu(300 * px), Emu(120 * px), Emu(60 * px))
    for box in (a, b):
        box.fill.solid()
    # the exporter's rendering of <path d="M220 130 H310 V330 H400" marker-end=...>: one open custom-geometry freeform with a tail end
    w, h = 180 * px, 200 * px
    xml = (f'<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
           f'<p:nvSpPr><p:cNvPr id="50" name="Freeform 50"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr>'
           f'<a:xfrm><a:off x="{220 * px}" y="{130 * px}"/><a:ext cx="{w}" cy="{h}"/></a:xfrm>'
           f'<a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/><a:rect l="l" t="t" r="r" b="b"/><a:pathLst><a:path w="{w}" h="{h}">'
           f'<a:moveTo><a:pt x="0" y="0"/></a:moveTo><a:lnTo><a:pt x="{90 * px}" y="0"/></a:lnTo><a:lnTo><a:pt x="{90 * px}" y="{h}"/></a:lnTo>'
           f'<a:lnTo><a:pt x="{w}" y="{h}"/></a:lnTo></a:path></a:pathLst></a:custGeom><a:noFill/>'
           f'<a:ln w="14288"><a:solidFill><a:srgbClr val="1F2937"/></a:solidFill><a:tailEnd type="triangle"/></a:ln></p:spPr></p:sp>')
    slide.shapes._spTree.append(etree.fromstring(xml))
    prs.save(path)
    return a.shape_id, b.shape_id


def test_an_exported_elbow_arrow_becomes_a_connector_glued_to_both_boxes(tmp_path):
    import zipfile
    import re
    import pptx_text_in_shapes
    source, target = tmp_path / "in.pptx", tmp_path / "out.pptx"
    first, second = _deck_with_elbow(source)
    pptx_text_in_shapes.convert(source, target)
    xml = zipfile.ZipFile(target).read("ppt/slides/slide1.xml").decode("utf-8")
    connector = re.search(r"<p:cxnSp>.*?</p:cxnSp>", xml, re.S).group(0)
    assert 'prst="bentConnector3"' in connector and 'name="adj1" fmla="val 50000"' in connector
    assert f'<a:stCxn id="{first}" idx="3"/>' in connector and f'<a:endCxn id="{second}" idx="1"/>' in connector
    assert "tailEnd" in connector and "custGeom" not in connector
    off = re.search(r'<a:off x="(\d+)" y="(\d+)"/>\s*<a:ext cx="(\d+)" cy="(\d+)"/>', connector)
    assert tuple(int(v) // 9525 for v in off.groups()) == (220, 130, 180, 200)  # nothing moved
