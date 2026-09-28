"""deck_consistency.py (failure family F06): cross-page drift found on synthetic page SVGs, and its place in the deck runner."""

import argparse
import importlib
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import deck_consistency as dc  # noqa: E402

HEAD = '<svg xmlns="http://www.w3.org/2000/svg" width="1280" height="720" viewBox="0 0 1280 720">'


def _chrome(number: int, folio: str | None = None, rule: bool = True, footer_fill: str = "#5F6368") -> str:
    folio = str(number) if folio is None else folio
    return ('<g id="chrome" data-pptx-role="chrome">'
            '<rect x="0" y="0" width="1280" height="720" fill="#FFFFFF"/>'
            '<text x="54" y="35" font-size="10" fill="#202522">Meridian Life Assurance</text>'
            + ('<line x1="54" y1="676" x2="1226" y2="676" stroke="#D8D5CE"/>' if rule else "")
            + f'<text x="760" y="696" font-size="10" fill="{footer_fill}">MAPH proposal · Confidential</text>'
            f'<text x="1226" y="696" font-size="10" text-anchor="end" fill="#202522">{folio}</text></g>')


def _page(number: int, body: str, **chrome) -> str:
    return HEAD + _chrome(number, **chrome) + body + "</svg>"


def _text(x: int, y: int, words: str, size: int = 16, eid: str = "", **attrs) -> str:
    extra = "".join(f' {k.replace("_", "-")}="{v}"' for k, v in attrs.items())
    return f'<text{f" id={chr(34)}{eid}{chr(34)}" if eid else ""} x="{x}" y="{y}" font-size="{size}"{extra}>{words}</text>'


def _deck(tmp_path: Path, pages: dict[int, str], spec: str = "") -> Path:
    project = tmp_path / "deck"
    (project / "svg_output").mkdir(parents=True)
    for number, svg in pages.items():
        (project / "svg_output" / f"{number:02d}_page.svg").write_text(svg, encoding="utf-8")
    if spec:
        (project / "design_spec.md").write_text(spec, encoding="utf-8")
    return project


def _body_pages(extra: dict[int, str], count: int = 5, **chrome_by_page) -> dict[int, str]:
    pages = {1: HEAD + '<rect x="0" y="0" width="1280" height="720" fill="#0B5D48"/>' + _text(80, 300, "Cover", 40) + "</svg>"}
    for number in range(2, count + 1):
        pages[number] = _page(number, extra.get(number, _text(80, 200, f"Body of page {number}")), **chrome_by_page.get(number, {}))
    return pages


def _kinds(report: dict, kind: str) -> list[dict]:
    return [f for f in report["findings"] if f["kind"] == kind]


# --- figures ------------------------------------------------------------------------------------------------------------------

def test_one_fee_with_two_values_on_two_pages_is_a_certain_conflict_addressed_to_the_page_that_leaves_the_plan(tmp_path):
    spec = "## VIII\nThe all-in maximum is USD 98,920 (USD 94,120 base + USD 4,800 contingency).\n"
    project = _deck(tmp_path, _body_pages({
        2: _text(80, 200, "Our offer: USD 94,120 base + up to USD 4,800 contingency = USD 98,920 maximum"),
        3: _text(80, 200, "Fee equation: 94,120 base + 4,800 contingency + 0 other costs = USD 98,720 maximum")}), spec)
    report = dc.check(project, project / "design_spec.md")
    conflicts = _kinds(report, "FIGURE_CONFLICT")
    assert len(conflicts) == 1 and conflicts[0]["severity"] == "certain" and conflicts[0]["pages"] == [2, 3]
    assert "98,720" in conflicts[0]["message"] and "98,920" in conflicts[0]["message"]
    assert list(conflicts[0]["repair"]) == ["3"]  # only P03 departs from the plan's figure
    assert dc.repair_items(report) == {3: [f"[{conflicts[0]['id']} FIGURE_CONFLICT] " + conflicts[0]["repair"]["3"]]}


def test_a_drawn_table_row_and_a_summary_line_that_disagree_conflict_but_baseline_and_target_do_not(tmp_path):
    summary = _text(80, 200, "312 incidents a year; 4 h 50 min mean time to resolve (MTTR); 58% are regressions")
    table = (_text(80, 180, "MEASURE", 12, font_weight="bold") + _text(360, 180, "BASELINE", 12, font_weight="bold")
             + _text(620, 180, "YEAR-ONE TARGET", 12, font_weight="bold")
             + _text(80, 240, "Mean time to resolve P1/P2") + _text(360, 240, "{baseline}") + _text(620, 240, "2 h 25 min (−50%)"))
    consistent = _deck(tmp_path / "a", _body_pages({2: summary, 4: table.replace("{baseline}", "4 h 50 min")}))
    assert _kinds(dc.check(consistent), "FIGURE_CONFLICT") == []  # 2 h 25 min is the target, not a second baseline
    drifted = _deck(tmp_path / "b", _body_pages({2: summary, 4: table.replace("{baseline}", "4 h 30 min")}))
    conflicts = _kinds(dc.check(drifted), "FIGURE_CONFLICT")
    assert len(conflicts) == 1 and "4 h 50 min" in conflicts[0]["message"] and "4 h 30 min" in conflicts[0]["message"]
    assert conflicts[0]["pages"] == [2, 4] and set(conflicts[0]["repair"]) == {"2", "4"}  # no plan figure: both pages reconcile


def test_the_same_duration_written_with_or_without_min_is_one_value(tmp_path):
    project = _deck(tmp_path, _body_pages({2: _text(80, 200, "Today: 4 h 50 min mean time to resolve"),
                                           3: _text(80, 200, "Baseline 4 h 50 mean time to resolve")}))
    report = dc.check(project)
    assert _kinds(report, "FIGURE_CONFLICT") == []
    drift = _kinds(report, "FORMAT_DRIFT")
    assert len(drift) == 1 and drift[0]["severity"] == "flagged"


def test_a_numbered_pair_gives_a_before_and_after_transition_its_runs(tmp_path):
    chart = ('<g id="scores" data-pptx-replace-with="chart"><metadata type="application/json">'
             + json.dumps({"x": 600, "y": 100, "width": 600, "height": 400, "type": "bar", "categories": ["Logic"],
                           "series": [{"name": "kirkland1", "values": [7]}, {"name": "kirkland2", "values": [8]}]})
             + "</metadata></g>")
    table = ('<g id="run-costs" data-pptx-replace-with="table"><metadata type="application/json">'
             + json.dumps({"x": 80, "y": 100, "width": 900, "height": 200,
                           "columns": [{"text": "Run"}, {"text": "Pages"}, {"text": "Model fee"}],
                           "rows": [[{"text": "kirkland1"}, {"text": "12"}, {"text": "{fee}"}],
                                    [{"text": "kirkland2"}, {"text": "12"}, {"text": "$9.14 plan→check"}]]})
             + "</metadata></g>")
    story = _text(80, 200, "Panel: not shortlisted → shortlisted; one unresolved timeline; $5.24 → $9.14.") + chart
    drifted = _deck(tmp_path / "a", _body_pages({2: story, 3: table.replace("{fee}", "$5.49 whole")}))
    conflicts = _kinds(dc.check(drifted), "FIGURE_CONFLICT")
    assert len(conflicts) == 1 and "$5.24" in conflicts[0]["message"] and "$5.49" in conflicts[0]["message"]
    assert "kirkland1" in conflicts[0]["fix"]
    same = _deck(tmp_path / "b", _body_pages({2: story, 3: table.replace("{fee}", "$5.24")}))
    assert _kinds(dc.check(same), "FIGURE_CONFLICT") == []  # chart scores (7, 8) are not the page counts (12) either


def test_different_line_items_a_digit_apart_in_tables_and_labelled_states_are_not_conflicts(tmp_path):
    project = _deck(tmp_path, _body_pages({
        2: _text(80, 200, "Change-failure rate 12% → 7% in year one"),
        3: _text(80, 200, "12% change-failure rate vs the 5% target"),
        4: _text(80, 200, "Delivery Manager") + _text(400, 200, "65 days"),
        5: _text(80, 200, "Manager") + _text(400, 200, "75 days")}))
    assert _kinds(dc.check(project), "FIGURE_CONFLICT") == []


# --- formats ------------------------------------------------------------------------------------------------------------------

def test_one_fee_written_as_millions_and_in_full_is_format_drift_but_weeks_and_months_are_not(tmp_path):
    project = _deck(tmp_path, _body_pages({2: _text(80, 200, "A traceable path in 26 weeks for a fixed HKD 7.507m"),
                                           3: _text(80, 200, "Fixed professional fee: HKD 7,507,000, exclusive of taxes"),
                                           4: _text(80, 200, "Six months from mobilisation to adoption")}))
    report = dc.check(project)
    drift = _kinds(report, "FORMAT_DRIFT")
    assert len(drift) == 1 and drift[0]["pages"] == [2, 3] and "7.507m" in drift[0]["message"] and "7,507,000" in drift[0]["message"]
    assert _kinds(report, "FIGURE_CONFLICT") == []


# --- terms --------------------------------------------------------------------------------------------------------------------

def test_suffix_variants_of_one_name_are_term_drift_and_plurals_are_not(tmp_path):
    project = _deck(tmp_path, _body_pages({2: _text(80, 200, "The Review Agent checks every PR before a human approves it."),
                                           3: _text(80, 200, "The Reviewer agent posts its verdict to the PR."),
                                           4: _text(80, 200, "Both Decision Boards meet monthly; the Decision Board chair decides.")}))
    drift = _kinds(dc.check(project), "TERM_DRIFT")
    assert len(drift) == 1 and drift[0]["severity"] == "flagged" and drift[0]["pages"] == [2, 3]
    assert "Review Agent" in drift[0]["message"] and "Reviewer agent" in drift[0]["message"]


def test_a_step_name_stacked_over_its_owner_reads_as_one_name(tmp_path):
    strip = '<g id="loop">' + _text(80, 300, "Diagnose", 14) + _text(80, 318, "agent", 11) + "</g>"
    project = _deck(tmp_path, _body_pages({2: _text(80, 200, "The Diagnosis Agent reads traces and code."),
                                           3: _text(80, 200, "The Diagnosis Agent never writes."), 4: strip}))
    drift = _kinds(dc.check(project), "TERM_DRIFT")
    assert len(drift) == 1 and "'Diagnose agent' on P04" in drift[0]["message"] and "P04" in drift[0]["fix"]


def test_a_glossary_in_the_plan_names_the_canonical_form(tmp_path):
    spec = "# Spec\n\n## Glossary\n\n- Remediation Agent: proposes the fix (Propose step).\n- **Gate 1**: pilot decision.\n\n## IX\n"
    project = _deck(tmp_path, _body_pages({2: _text(80, 200, "The Remediation Agent opens a branch."),
                                           3: _text(80, 200, "THE REMEDIATION AGENT NEVER MERGES"),
                                           4: _text(80, 200, "A remediation-agent PR needs two approvals.")}), spec)
    assert dc.glossary_terms(spec) == ["Remediation Agent", "Gate 1"]
    drift = _kinds(dc.check(project, project / "design_spec.md"), "TERM_DRIFT")
    assert len(drift) == 1 and drift[0]["pages"] == [4] and "'Remediation Agent'" in drift[0]["fix"]  # the caps label is typography


# --- chrome and markers -------------------------------------------------------------------------------------------------------

def test_a_wrong_folio_and_a_missing_footer_rule_are_certain_and_a_recoloured_footer_is_flagged(tmp_path):
    pages = _body_pages({}, count=7)
    pages[4] = _page(4, _text(80, 200, "Body"), folio="5")
    pages[5] = _page(5, _text(80, 200, "Body"), rule=False)
    pages[6] = _page(6, _text(80, 200, "Body"), footer_fill="#B4162E")
    project = _deck(tmp_path, pages)
    chrome = _kinds(dc.check(project), "CHROME_DRIFT")
    certain = [f for f in chrome if f["severity"] == "certain"]
    assert sorted((f["pages"][0], "folio" in f["message"]) for f in certain) == [(4, True), (5, False)]
    assert "says 5" in next(f for f in certain if f["pages"] == [4])["message"]
    flagged = [f for f in chrome if f["severity"] == "flagged"]
    assert [f["pages"] for f in flagged] == [[6]] and "#b4162e" in flagged[0]["message"]


def test_section_trackers_and_source_lines_that_change_per_page_are_not_chrome_drift(tmp_path):
    pages = {n: _page(n, _text(1226, 35, f"SECTION {'AB'[n % 2]}", 10, text_anchor="end") + _text(54, 690, f"Source: note {n * 7}", 10))
             for n in range(2, 7)}
    project = _deck(tmp_path, {1: HEAD + "</svg>", **pages})
    assert _kinds(dc.check(project), "CHROME_DRIFT") == []


def test_one_label_keyed_by_different_marker_fills_is_marker_drift(tmp_path):
    def gate(fill: str) -> str:
        return (f'<rect x="831" y="474" width="12" height="12" fill="{fill}" transform="rotate(45 837 480)"/>'
                + _text(858, 485, "Gate 2 (week 24)", 14))
    project = _deck(tmp_path, _body_pages({2: gate("#4F6D8A"), 3: gate("#B4162E"), 4: gate("#B4162E")}))
    drift = _kinds(dc.check(project), "MARKER_DRIFT")
    assert len(drift) == 1 and drift[0]["pages"] == [2, 3, 4] and list(drift[0]["repair"]) == ["2"]


# --- inputs and reports -------------------------------------------------------------------------------------------------------

def test_an_exported_pptx_is_read_slide_by_slide(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    deck = Presentation()
    for words in ("Run cost for kirkland1 whole run", "Fee summary"):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = words
    rows = deck.slides[0].shapes.add_table(2, 2, Inches(1), Inches(3), Inches(6), Inches(1)).table
    rows.cell(0, 0).text, rows.cell(0, 1).text, rows.cell(1, 0).text, rows.cell(1, 1).text = "Run", "Model fee", "kirkland1", "$5.49"
    deck.slides[1].shapes.add_textbox(Inches(1), Inches(3), Inches(6), Inches(1)).text_frame.text = "Model fee for kirkland1: $5.24"
    path = tmp_path / "deck.pptx"
    deck.save(str(path))
    report = dc.check(path)
    assert report["input"] == "pptx" and len(report["pages"]) == 2
    conflicts = _kinds(report, "FIGURE_CONFLICT")
    assert len(conflicts) == 1 and conflicts[0]["pages"] == [1, 2]


def test_the_cli_writes_its_reports_into_the_project_review_folder(tmp_path, monkeypatch, capsys):
    project = _deck(tmp_path, _body_pages({2: _text(80, 200, "Fee HKD 7.507m"), 3: _text(80, 200, "Fee HKD 7,507,000")}))
    monkeypatch.setattr(sys, "argv", ["deck_consistency.py", str(project)])
    assert dc.main() == 0
    written = json.loads((project / ".review" / "consistency.json").read_text(encoding="utf-8"))
    assert written["flagged"] == 1 and (project / ".review" / "consistency.md").is_file()
    assert "FORMAT_DRIFT" in capsys.readouterr().out
    block = dc.prompt_block(written)
    assert block.startswith("DETERMINISTIC CROSS-PAGE FINDINGS TO CONFIRM") and "C01 FLAGGED FORMAT_DRIFT" in block


# --- the deck runner ----------------------------------------------------------------------------------------------------------

def _runner_module():
    host_dir = Path(__file__).resolve().parents[4] / "hosts" / "responses_api"
    if str(host_dir) not in sys.path:
        sys.path.insert(0, str(host_dir))
    return importlib.import_module("deck_runner")


def test_the_runner_sends_certain_findings_to_their_pages_and_records_what_is_left(tmp_path):
    runner = _runner_module()
    spec = "The all-in maximum is USD 98,920.\n"
    project = _deck(tmp_path, _body_pages({2: _text(80, 200, "Our offer: USD 98,920 maximum"),
                                           3: _text(80, 200, "Fee equation = USD 98,720 maximum"),
                                           4: _text(80, 200, "Fee HKD 7.507m"), 5: _text(80, 200, "Fee HKD 7,507,000")}), spec)
    bare = object.__new__(runner.Runner)
    bare.args = argparse.Namespace(repair_rounds=2, max_parallel=2)
    bare.project, bare.page_sessions, bare.authored_by = project, {"02_page": ["s2"], "03_page": ["s3"], "04_page": ["s4"]}, {}
    bare.consistency_before, bare.consistency_after = {}, {}
    said, repaired = [], []
    bare.say = said.append
    bare.repair = lambda page, issues, hint="": repaired.append((page["stem"], issues, hint))
    pages = [{"number": n, "stem": f"{n:02d}_page"} for n in range(1, 6)]
    bare.consistency_before = bare.consistency("before review")
    assert (project / ".review" / "consistency.md").is_file() and bare.consistency_before["certain"] == 1
    assert any(line.startswith("consistency (before review): 1 certain, 1 flagged") for line in said)
    bare.deck_repair(pages)  # no deck review on file: the measured findings still go to their page
    assert [(stem, len(items)) for stem, items, _ in repaired] == [("03_page", 1)]  # the flagged format drift is the reviewer's to confirm
    assert "98,720" in repaired[0][1][0] and "measured by a script" in repaired[0][2]
    (project / "svg_output" / "03_page.svg").write_text(_page(3, _text(80, 200, "Fee equation = USD 98,920 maximum")), encoding="utf-8")
    bare.consistency_after = bare.consistency("after deck repair")
    assert (project / ".review" / "consistency_after_repair.md").is_file()
    lines = bare.consistency_lines()
    assert "Before the deck review: 1 certain, 1 flagged" in lines[0] and "After deck repair: 0 certain, 1 flagged" in lines[0]
    assert len(lines) == 1  # nothing certain left to list


def test_the_deck_review_is_given_the_findings_to_confirm(tmp_path, monkeypatch):
    runner = _runner_module()
    import page_review
    monkeypatch.setattr(runner, "ROOT", tmp_path)  # the runner names its outputs relative to the checkout
    project = _deck(tmp_path, _body_pages({2: _text(80, 200, "Fee HKD 7.507m"), 3: _text(80, 200, "Fee HKD 7,507,000")}))
    (project / ".preview").mkdir()
    (project / ".preview" / "contact_sheet.png").write_bytes(b"png")
    bare = object.__new__(runner.Runner)
    bare.project, bare.say = project, (lambda text: None)
    bare.reviewers = {"frontier": {"model": "m", "effort": "high"}}
    bare.reviewer_env = lambda tier: {}
    bare.script = lambda *a, **k: argparse.Namespace(stdout="")
    bare.consistency_before = dc.check(project)
    sent = {}
    monkeypatch.setattr(page_review, "_image_item", lambda image: [{"type": "input_image", "image_url": "data:"}])
    monkeypatch.setattr(page_review, "_review_call", lambda payload: (sent.update(payload) or "DECK VERDICT: PASS", {}))
    monkeypatch.setattr(page_review, "update_journal", lambda *a, **k: None)
    bare.deck_review([{"number": n, "stem": f"{n:02d}_page"} for n in range(1, 6)])
    texts = [part["text"] for part in sent["input"][0]["content"] if part["type"] == "input_text"]
    assert any(text.startswith("DETERMINISTIC CROSS-PAGE FINDINGS TO CONFIRM") and "FORMAT_DRIFT" in text for text in texts)
    assert "Confirm each one you can see in the transcript" in sent["instructions"]
