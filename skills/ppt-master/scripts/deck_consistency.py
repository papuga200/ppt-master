#!/usr/bin/env python3
"""Deterministic cross-page consistency check for a deck (failure family F06).

Reads every page of a deck - a project's ``svg_output/*.svg`` (or any folder of page SVGs), or an exported ``.pptx`` - and
reports what drifts between pages:

- ``FIGURE_CONFLICT`` (certain): one quantity stated with different values on different pages. A figure's context key is the
  label around it (the words before or after it in its clause, a table's row and column headers, a chart's category and
  series, and the numbered pair a ``A -> B`` transition refers to). Two pages conflict when their figures share a key and a
  unit but no value in common.
- ``FORMAT_DRIFT`` (flagged): the same value written in different formats (``HKD 7.507m`` against ``HKD 7,507,000``;
  ``4 h 50 min`` against ``290 min``). Different units that are not the same value (``26 weeks`` / ``6 months``) are not drift.
- ``TERM_DRIFT`` (flagged): variants of one named term (``Diagnose Agent`` / ``Diagnosis Agent``; ``Review Agent`` /
  ``Reviewer agent``). With a ``## Glossary`` (or ``## Canonical terms``) section in design_spec.md written as
  ``- Term: definition`` lines, any variant of a listed term is reported against the canonical form.
- ``CHROME_DRIFT`` (certain when a chrome element is missing or extra, or the folio is out of sequence; flagged for a style
  or wording difference): body pages compared with the majority, after normalising folio numbers and section-tracker text.
- ``MARKER_DRIFT`` (flagged): one short label drawn with a marker shape of a different fill on different pages.

Usage:
    deck_consistency.py <project | svg-folder | deck.pptx> [--spec design_spec.md] [--json OUT.json] [--md OUT.md]

With a project folder, ``--spec`` defaults to its ``design_spec.md`` and the reports go to ``.review/consistency.{json,md}``
unless ``--json``/``--md`` say otherwise. Exit code 0 always (the report is advisory); ``--fail-on-certain`` exits 1 when a
certain finding is open.
"""
from __future__ import annotations

import argparse
import json
import math
import re
import sys
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

CERTAIN, FLAGGED = "certain", "flagged"
SVG_NS = "{http://www.w3.org/2000/svg}"

# --------------------------------------------------------------------------------------------------------------------------
# page model
# --------------------------------------------------------------------------------------------------------------------------


@dataclass
class Item:
    """One run of drawn text (a <text>, a table cell, a chart value, a PowerPoint text frame)."""
    page: int
    eid: str
    text: str
    x: float = 0.0
    y: float = 0.0
    w: float = 0.0
    h: float = 0.0
    size: float = 0.0
    fill: str = ""
    weight: str = ""
    chrome: bool = False
    source: str = "text"  # text | table | chart
    row: str = ""
    col: str = ""
    col_unit: str = ""
    ax: float = 0.0        # the x the text is aligned on (its anchor), exact where the width is only estimated
    anchor: str = "start"
    band: bool = False     # chrome only because it sits in the header or footer band, not because a group says so
    layer: str = ""        # data-pptx-layer: master / layout elements come from the template


@dataclass
class Shape:
    page: int
    eid: str
    tag: str
    x: float
    y: float
    w: float
    h: float
    fill: str = ""
    stroke: str = ""
    chrome: bool = False
    layer: str = ""


@dataclass
class Page:
    number: int
    stem: str
    items: list[Item] = field(default_factory=list)
    shapes: list[Shape] = field(default_factory=list)
    master: str = ""
    layout: str = ""
    height: float = 720.0
    width: float = 1280.0


def _num(value, default: float = 0.0) -> float:
    if value is None:
        return default
    match = re.match(r"\s*(-?\d+(?:\.\d+)?)", str(value))
    return float(match.group(1)) if match else default


def _style(element: ET.Element) -> dict[str, str]:
    out = {}
    for part in (element.get("style") or "").split(";"):
        if ":" in part:
            key, value = part.split(":", 1)
            out[key.strip()] = value.strip()
    return out


def _matrix(transform: str) -> tuple[float, float, float, float, float, float]:
    """Affine matrix of a transform attribute (translate, scale, rotate, matrix; a skew is ignored)."""
    a, b, c, d, e, f = 1.0, 0.0, 0.0, 1.0, 0.0, 0.0
    for name, args in re.findall(r"(\w+)\s*\(([^)]*)\)", transform or ""):
        values = [float(v) for v in re.findall(r"-?\d+(?:\.\d+)?(?:e-?\d+)?", args)]
        if name == "translate" and values:
            m = (1, 0, 0, 1, values[0], values[1] if len(values) > 1 else 0.0)
        elif name == "scale" and values:
            m = (values[0], 0, 0, values[1] if len(values) > 1 else values[0], 0, 0)
        elif name == "rotate" and values:
            cos, sin = math.cos(math.radians(values[0])), math.sin(math.radians(values[0]))
            cx, cy = (values[1], values[2]) if len(values) == 3 else (0.0, 0.0)
            m = (cos, sin, -sin, cos, cx - cos * cx + sin * cy, cy - sin * cx - cos * cy)
        elif name == "matrix" and len(values) == 6:
            m = tuple(values)
        else:
            continue
        a, b, c, d, e, f = (a * m[0] + c * m[1], b * m[0] + d * m[1], a * m[2] + c * m[3], b * m[2] + d * m[3],
                            a * m[4] + c * m[5] + e, b * m[4] + d * m[5] + f)
    return a, b, c, d, e, f


def _compose(outer, inner):
    a, b, c, d, e, f = outer
    m = inner
    return (a * m[0] + c * m[1], b * m[0] + d * m[1], a * m[2] + c * m[3], b * m[2] + d * m[3],
            a * m[4] + c * m[5] + e, b * m[4] + d * m[5] + f)


def _apply(m, x: float, y: float) -> tuple[float, float]:
    return m[0] * x + m[2] * y + m[4], m[1] * x + m[3] * y + m[5]


def _text_of(element: ET.Element) -> str:
    """The words of a <text>: a <tspan> that starts a new line (own x or dy) is joined with a space, an inline run is not;
    an empty line is a paragraph break; a line ending in a hyphen joins the next without a space."""
    pieces: list[str] = [element.text or ""]
    for child in element:
        new_line = child.get("x") is not None or _num(child.get("dy")) > 0 or child.get("y") is not None
        inner = "".join(child.itertext())
        if new_line:
            previous = "".join(pieces).rstrip()
            if not inner.strip():
                pieces.append("\n")
            elif previous.endswith("-") and not previous.endswith(" -"):
                pieces[-1] = pieces[-1].rstrip()
                pieces.append(inner)
            else:
                pieces.append(" " + inner)
        else:
            pieces.append(inner)
        pieces.append(child.tail or "")
    text = "".join(pieces)
    return "\n".join(" ".join(line.split()) for line in text.split("\n")).strip()


CHROME_ROLES = {"chrome", "header", "footer", "logo", "background", "folio", "page-number", "slide-number"}


def _is_chrome_group(element: ET.Element) -> bool:
    eid = (element.get("id") or "").lower()
    role = (element.get("data-pptx-role") or "").lower()
    layer = (element.get("data-pptx-layer") or "").lower()
    return eid == "chrome" or role in CHROME_ROLES or layer in {"master", "layout"}


def _page_number(stem: str, index: int) -> int:
    match = re.match(r"^(?:[Pp]|slide[-_]?)?(\d{1,3})(?:\D|$)", stem)
    return int(match.group(1)) if match else index


def _shape_box(element: ET.Element, tag: str) -> tuple[float, float, float, float] | None:
    g = element.get
    if tag == "rect":
        return _num(g("x")), _num(g("y")), _num(g("width")), _num(g("height"))
    if tag == "circle":
        r = _num(g("r"))
        return _num(g("cx")) - r, _num(g("cy")) - r, 2 * r, 2 * r
    if tag == "ellipse":
        rx, ry = _num(g("rx")), _num(g("ry"))
        return _num(g("cx")) - rx, _num(g("cy")) - ry, 2 * rx, 2 * ry
    if tag == "line":
        x1, y1, x2, y2 = (_num(g(k)) for k in ("x1", "y1", "x2", "y2"))
        return min(x1, x2), min(y1, y2), abs(x2 - x1), abs(y2 - y1)
    if tag in {"polygon", "polyline"}:
        values = [float(v) for v in re.findall(r"-?\d+(?:\.\d+)?", g("points") or "")]
    elif tag == "path":
        d = g("d") or ""
        if re.search(r"[a-y]", d.replace("e", "")):  # relative commands: the numbers are offsets, not points
            first = re.findall(r"-?\d+(?:\.\d+)?", d)[:2]
            return (float(first[0]), float(first[1]), 0.0, 0.0) if len(first) == 2 else None
        values = [float(v) for v in re.findall(r"-?\d+(?:\.\d+)?", re.sub(r"[AH V]", " ", d))]
    else:
        return None
    xs, ys = values[0::2], values[1::2]
    if not xs or not ys:
        return None
    return min(xs), min(ys), max(xs) - min(xs), max(ys) - min(ys)


def load_svg_page(path: Path, index: int) -> Page:
    root = ET.parse(path).getroot()
    view = [float(v) for v in re.findall(r"-?\d+(?:\.\d+)?", root.get("viewBox") or "")]
    page = Page(number=_page_number(path.stem, index), stem=path.stem, master=root.get("data-pptx-master") or "",
                layout=root.get("data-pptx-layout") or "",
                width=view[2] if len(view) == 4 else _num(root.get("width"), 1280.0),
                height=view[3] if len(view) == 4 else _num(root.get("height"), 720.0))

    def walk(element: ET.Element, matrix, inherited: dict[str, str], chrome: bool, eid: str, layer: str) -> None:
        tag = element.tag.replace(SVG_NS, "")
        if tag in {"defs", "metadata", "style", "title", "desc", "clipPath", "mask", "symbol", "marker"}:
            return
        matrix = _compose(matrix, _matrix(element.get("transform") or ""))
        chrome = chrome or _is_chrome_group(element)
        eid = element.get("id") or eid
        layer = element.get("data-pptx-layer") or layer
        state = dict(inherited)
        for key in ("font-size", "fill", "font-weight", "text-anchor", "stroke"):
            value = element.get(key) or _style(element).get(key)
            if value:
                state[key] = value
        replace = element.get("data-pptx-replace-with")
        if replace in {"table", "chart"}:
            meta = next((child for child in element if child.tag.replace(SVG_NS, "") == "metadata"), None)
            if meta is not None and (meta.text or "").strip():
                try:
                    data = json.loads(meta.text)
                except ValueError:
                    data = None
                if isinstance(data, dict):
                    ox, oy = _apply(matrix, _num(data.get("x")), _num(data.get("y")))
                    box = (ox, oy, _num(data.get("width")), _num(data.get("height")))
                    page.items.extend(_table_items(page.number, eid, data, box) if replace == "table" else
                                      _chart_items(page.number, eid, data, box))
                    return  # the drawn fallback repeats the object's words
        if tag == "text":
            text = _text_of(element)
            if text:
                size = _num(state.get("font-size"), 16.0) * math.sqrt(abs(matrix[0] * matrix[3] - matrix[1] * matrix[2]) or 1.0)
                first = next((child for child in element if child.tag.replace(SVG_NS, "") == "tspan"), None)
                x0 = _num(element.get("x"), _num(first.get("x")) if first is not None else 0.0)
                y0 = _num(element.get("y"), _num(first.get("y")) if first is not None else 0.0)
                x, y = _apply(matrix, x0, y0)
                lines = [line for line in text.split("\n")] or [text]
                longest = max((len(line) for line in lines), default=len(text))
                line_count = max(1, 1 + sum(1 for child in element if child.tag.replace(SVG_NS, "") == "tspan"
                                            and (_num(child.get("dy")) > 0 or child.get("y") is not None)))
                width = longest * size * 0.52 / max(1, line_count) if line_count > 1 else len(text) * size * 0.52
                anchor = state.get("text-anchor", "start")
                left = x - width if anchor == "end" else x - width / 2 if anchor == "middle" else x
                page.items.append(Item(page=page.number, eid=eid or f"text@{x:.0f},{y:.0f}", text=text, x=left, y=y - size * 0.8,
                                       w=width, h=size * 1.2 * line_count, size=round(size, 1), fill=state.get("fill", "").lower(),
                                       weight=state.get("font-weight", ""), chrome=chrome, ax=x, anchor=anchor, layer=layer))
            return
        box = _shape_box(element, tag)
        if box is not None:
            corners = [_apply(matrix, box[0] + dx, box[1] + dy) for dx in (0.0, box[2]) for dy in (0.0, box[3])]
            x, y = min(c[0] for c in corners), min(c[1] for c in corners)
            width, height = max(c[0] for c in corners) - x, max(c[1] for c in corners) - y
            page.shapes.append(Shape(page=page.number, eid=eid or f"{tag}@{x:.0f},{y:.0f}", tag=tag, x=x, y=y, w=width,
                                     h=height, fill=(element.get("fill") or _style(element).get("fill") or state.get("fill", "")).lower(),
                                     stroke=(element.get("stroke") or _style(element).get("stroke") or "").lower(), chrome=chrome, layer=layer))
        for child in element:
            walk(child, matrix, state, chrome, eid, layer)

    walk(root, (1.0, 0.0, 0.0, 1.0, 0.0, 0.0), {}, False, "", "")
    for item in page.items:  # a text in the header or footer band is chrome whatever group holds it
        if item.source == "text" and not item.chrome and (item.y + item.h <= page.height * 0.085 or item.y >= page.height * 0.915):
            item.chrome = item.band = True
    return page


def _cell_text(cell) -> str:
    """A table cell's words, whatever shape its JSON takes (text, paragraphs, runs)."""
    def flat(value) -> str:
        if isinstance(value, dict):
            for key in ("paragraphs", "runs", "text"):
                if key in value:
                    return flat(value[key])
            return ""
        if isinstance(value, list):
            return " ".join(flat(v) for v in value)
        return str(value if value is not None else "")
    return " ".join(flat(cell).split())


def _table_items(number: int, eid: str, data: dict, box) -> list[Item]:
    columns = [_cell_text(c) for c in data.get("columns") or []]
    rows = data.get("rows") or []
    heights = data.get("row_heights") or []
    items: list[Item] = []
    y = box[1] + (_num(heights[0]) if heights else 30.0)
    for r, row in enumerate(rows):
        cells = [_cell_text(c) for c in row]
        header = cells[0] if cells and not _is_numeric_text(cells[0]) else ""
        for j, text in enumerate(cells):
            if not text:
                continue
            col = columns[j] if j < len(columns) else ""
            items.append(Item(page=number, eid=f"{eid}[r{r + 1}c{j + 1}]", text=text, x=box[0], y=y, w=box[2], h=18.0,
                              source="table", row=header if j else "", col=col, col_unit=_unit_from_header(col)))
        y += _num(heights[r + 1]) if r + 1 < len(heights) else 24.0
    return items


def _chart_items(number: int, eid: str, data: dict, box) -> list[Item]:
    categories = [str(c) for c in data.get("categories") or []]
    items: list[Item] = []
    for series in data.get("series") or []:
        for index, value in enumerate(series.get("values") or []):
            if value is None:
                continue
            items.append(Item(page=number, eid=f"{eid}[{series.get('name', '')}/{categories[index] if index < len(categories) else index}]",
                              text=_fmt_plain(value), x=box[0], y=box[1], w=box[2], h=box[3], source="chart",
                              row=categories[index] if index < len(categories) else "", col=str(series.get("name") or "")))
        items.append(Item(page=number, eid=f"{eid}[series]", text=str(series.get("name") or ""), x=box[0], y=box[1], source="chart"))
    for category in categories:
        items.append(Item(page=number, eid=f"{eid}[category]", text=category, x=box[0], y=box[1], source="chart"))
    return items


def _fmt_plain(value) -> str:
    return str(int(value)) if isinstance(value, float) and value.is_integer() else str(value)


def load_pptx_pages(path: Path) -> list[Page]:
    from pptx import Presentation
    from pptx.util import Emu

    deck = Presentation(str(path))
    scale = 1280.0 / float(deck.slide_width or Emu(12192000))
    height = float(deck.slide_height or Emu(6858000)) * scale
    pages: list[Page] = []

    def walk(page: Page, shapes, chrome_names: set[str]) -> None:
        for shape in shapes:
            x = float(shape.left or 0) * scale
            y = float(shape.top or 0) * scale
            w = float(shape.width or 0) * scale
            h = float(shape.height or 0) * scale
            name = shape.name or f"shape{shape.shape_id}"
            if shape.shape_type is not None and shape.shape_type == 6:  # group
                walk(page, shape.shapes, chrome_names)
                continue
            chrome = False
            if shape.is_placeholder:
                kind = str(shape.placeholder_format.type).lower()
                chrome = any(word in kind for word in ("footer", "slide_number", "date"))
            if getattr(shape, "has_table", False) and shape.has_table:
                table = shape.table
                columns = [" ".join(c.text.split()) for c in table.rows[0].cells] if len(table.rows) else []
                for r, row in enumerate(list(table.rows)[1:], start=1):
                    cells = [" ".join(c.text.split()) for c in row.cells]
                    header = cells[0] if cells and not _is_numeric_text(cells[0]) else ""
                    for j, text in enumerate(cells):
                        if text:
                            col = columns[j] if j < len(columns) else ""
                            page.items.append(Item(page=page.number, eid=f"{name}[r{r}c{j + 1}]", text=text, x=x, y=y, w=w, h=h,
                                                   source="table", row=header if j else "", col=col, col_unit=_unit_from_header(col)))
                for j, text in enumerate(columns):
                    if text:
                        page.items.append(Item(page=page.number, eid=f"{name}[r0c{j + 1}]", text=text, x=x, y=y, source="table"))
                continue
            if getattr(shape, "has_chart", False) and shape.has_chart:
                try:
                    plot = shape.chart.plots[0]
                    categories = [str(c) for c in plot.categories]
                    data = {"categories": categories, "series": [{"name": s.name, "values": list(s.values)} for s in plot.series]}
                    page.items.extend(_chart_items(page.number, name, data, (x, y, w, h)))
                except Exception:  # noqa: BLE001 - an unreadable chart is skipped, not fatal
                    pass
                continue
            if shape.has_text_frame and shape.text_frame.text.strip():
                text = "\n".join(" ".join(p.text.split()) for p in shape.text_frame.paragraphs).strip()
                sizes = [r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs if r.font.size is not None]
                page.items.append(Item(page=page.number, eid=name, text=text, x=x, y=y, w=w, h=h,
                                       size=round(max(sizes) * 96 / 72, 1) if sizes else 0.0, chrome=chrome))

    for index, slide in enumerate(deck.slides, start=1):
        page = Page(number=index, stem=f"slide{index:02d}", height=height, layout=slide.slide_layout.name or "")
        walk(page, slide.shapes, set())
        for item in page.items:
            if item.source == "text" and not item.chrome and (item.y + item.h <= height * 0.085 or item.y >= height * 0.915):
                item.chrome = item.band = True
        pages.append(page)
    return pages


def load_pages(target: Path) -> tuple[list[Page], str]:
    """Pages of a project, an SVG folder or a .pptx, and a label for the report."""
    if target.suffix.lower() == ".pptx":
        return load_pptx_pages(target), "pptx"
    folder = target / "svg_output" if (target / "svg_output").is_dir() else target
    files = sorted(p for p in folder.glob("*.svg"))
    return [load_svg_page(path, index) for index, path in enumerate(files, start=1)], "svg"


# --------------------------------------------------------------------------------------------------------------------------
# figures
# --------------------------------------------------------------------------------------------------------------------------

CURRENCIES = {"US$": "USD", "USD": "USD", "$": "USD", "HK$": "HKD", "HKD": "HKD", "S$": "SGD", "SGD": "SGD", "A$": "AUD", "AUD": "AUD",
              "C$": "CAD", "CAD": "CAD", "NZ$": "NZD", "NZD": "NZD", "EUR": "EUR", "€": "EUR", "GBP": "GBP", "£": "GBP", "CHF": "CHF",
              "JPY": "JPY", "¥": "JPY", "CNY": "CNY", "RMB": "CNY", "INR": "INR", "₹": "INR"}
CURRENCY_RE = re.compile(r"(US\$|HK\$|NZ\$|S\$|A\$|C\$|USD|HKD|SGD|AUD|CAD|NZD|EUR|GBP|CHF|JPY|CNY|RMB|INR|\$|€|£|¥|₹)\s?[~≈]?\s?$")
SCALES = {"k": 1e3, "K": 1e3, "thousand": 1e3, "m": 1e6, "M": 1e6, "mn": 1e6, "million": 1e6, "bn": 1e9, "b": 1e9, "B": 1e9, "billion": 1e9}
MONTHS = {"jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "sept", "oct", "nov", "dec", "january", "february", "march",
          "april", "june", "july", "august", "september", "october", "november", "december"}
ORDINAL_WORDS = {"week", "weeks", "wk", "w", "gate", "stage", "step", "phase", "wave", "p", "page", "pages", "slide", "slides", "§", "section",
                 "figure", "fig", "table", "v", "version", "q", "fy", "level", "tier", "option", "no", "#", "round", "rev", "run", "day",
                 "sprint", "release", "iteration", "item", "part", "chapter", "lane", "row", "column", "rank", "priority", "p1", "p2"}
STOP = {"a", "an", "the", "and", "or", "of", "to", "in", "on", "for", "per", "with", "at", "by", "from", "vs", "than", "is", "are", "was",
        "were", "be", "been", "it", "its", "this", "that", "these", "those", "as", "into", "over", "under", "up", "each", "every", "all",
        "no", "not", "one", "two", "three", "four", "five", "six", "seven", "eight", "nine", "ten", "about", "around", "approx", "c",
        "about", "some", "any", "only", "also", "which", "who", "what", "when", "where", "how", "our", "your", "their", "we", "you",
        "they", "he", "she", "has", "have", "had", "will", "would", "can", "could", "may", "might", "should", "must", "do", "does",
        "did", "so", "if", "but", "then", "there", "here", "more", "less", "most", "least", "same", "new", "via", "within", "across",
        "after", "before", "until", "while", "e.g", "i.e", "etc", "x", "×", "≈", "~", "see", "used", "use", "uses", "using", "up",
        "incl", "including", "excluding", "excl", "plus", "minus", "total", "totals"}
APPROX_BEFORE = re.compile(r"(?:~|≈|\babout|\bapprox\.?|\bapproximately|\baround|\bc\.|\broughly|\bnearly|\balmost|\bsome)\s*$", re.I)
BOUND_BEFORE = re.compile(r"(?:<|>|≤|≥|=<|>=|\bup to|\bunder|\bover|\bbelow|\babove|\bat least|\bat most|\bmax(?:imum)?|\bmin(?:imum)?|"
                          r"\bless than|\bmore than|\bfewer than|\bwithin|\bexceeds?|\bbeyond|\bceiling of|\bcap(?:ped)? (?:at|of))\s*$", re.I)
NUMBER_RE = re.compile(r"\d[\d,]*(?:\.\d+)?")


@dataclass
class Figure:
    page: int
    item: Item
    raw: str
    value: float
    unit: str            # USD, HKD, %, time, weeks, n:engineer, plain ...
    fmt: str             # the written form's class, for FORMAT_DRIFT
    decimals: int        # precision of the written value (for rounding tolerance)
    scale: float
    start: int
    end: int
    approx: bool = False
    bound: bool = False
    delta: bool = False
    ranged: bool = False
    key: frozenset = frozenset()
    anchors: frozenset = frozenset()
    rowkey: frozenset = frozenset()
    colkey: frozenset = frozenset()
    label: str = ""


def _is_numeric_text(text: str) -> bool:
    return bool(text) and not re.search(r"[A-Za-z]{3,}", text) and bool(re.search(r"\d", text))


UNIT_HEADERS = {"min": "time", "mins": "time", "minutes": "time", "h": "time", "hours": "time", "hrs": "time", "%": "%", "percent": "%",
                "weeks": "weeks", "days": "days", "months": "months", "years": "years", "usd": "USD", "hkd": "HKD", "sgd": "SGD",
                "eur": "EUR", "gbp": "GBP", "$": "USD"}


def _unit_from_header(header: str) -> str:
    """A column headed `Min`, `Hours`, `% share`, `(USD)` gives its bare numbers that unit."""
    words = re.findall(r"[a-z%$]+", header.lower())
    for word in words[-2:] + words[:1]:
        if word in UNIT_HEADERS:
            return UNIT_HEADERS[word] + ("*60" if word in {"h", "hours", "hrs"} else "")
    return ""


def _decimals(raw: str) -> int:
    return len(raw.split(".", 1)[1]) if "." in raw else 0


def extract_figures(item: Item) -> list[Figure]:
    """Every quantity in one text item, with its unit, written form and flags. Ordinals, dates, years, identifiers and
    ratios are not quantities and are skipped."""
    text = item.text
    figures: list[Figure] = []
    taken = 0  # the end of the last figure: `50` inside `4 h 50 min` is not a second figure
    for match in NUMBER_RE.finditer(text):
        raw = match.group(0).rstrip(",.")
        start, end = match.start(), match.start() + len(raw)
        if start < taken:
            continue
        if raw.count(",") and not re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", raw):
            raw = raw.split(",")[0]
            end = start + len(raw)
        before, after = text[:start], text[end:]
        if before and (before[-1].isalnum() or before[-1] in "._/:@#'"):
            continue  # part of an identifier, version, ratio or time of day (kirkland1, v3.4, 7/9, 10:30)
        if re.match(r"^(?:[/:]\d|\.\d|[A-Za-z]*\d)", after) or re.match(r"^[A-Za-z]{0,2}\d", after):
            continue
        if re.match(r"^\s*/\s*\d", after) or re.search(r"\d\s*/\s*$", before):
            continue
        try:
            value = float(raw.replace(",", ""))
        except ValueError:
            continue
        word_before = (re.findall(r"([A-Za-z§#]+)\.?\s*$", before) or [""])[0].lower()
        word_after = (re.findall(r"^\s*([A-Za-z]+)", after) or [""])[0].lower()
        if word_before in ORDINAL_WORDS or word_before in MONTHS or word_after in MONTHS:
            continue
        currency_match = CURRENCY_RE.search(before)
        currency = CURRENCIES.get(currency_match.group(1)) if currency_match else ""
        marker = currency_match.group(1) if currency_match else ""
        prefix_start = currency_match.start() if currency_match else start
        scale, fmt, unit, consumed = 1.0, "", "", 0
        decimals = _decimals(raw)
        if currency:
            m = re.match(r"^\s?(k|K|mn|m|M|bn|b|B|million|billion|thousand)\b", after)
            if m:
                scale, consumed = SCALES[m.group(1)], m.end()
                fmt = f"{marker}+{'compact' if len(m.group(1)) <= 2 else 'words'}"
            else:
                fmt = f"{marker}+{'grouped' if ',' in raw else 'digits'}"
            unit = currency
            value *= scale
        else:
            m = re.match(r"^\s?(?:%|per ?cent\b)", after, re.I)
            hm = re.match(r"^\s?h(?:ours?|rs?)?\s?(\d{1,2})(?:\s?(min(?:utes?|s)?|m)\b)?", after)
            if m:
                unit, consumed, fmt = "%", m.end(), "%" if "%" in m.group(0) else "per cent"
            elif re.match(r"^\s?(?:pp|percentage points?)\b", after):
                unit, consumed, fmt = "pp", re.match(r"^\s?(?:pp|percentage points?)\b", after).end(), "pp"
            elif hm and "." not in raw:
                value = value * 60 + float(hm.group(1))
                unit, consumed, fmt = "time", hm.end(), "h+min" if hm.group(2) else "h+bare"
                decimals = 0
            elif re.match(r"^\s?(?:min(?:utes?|s)?)\b", after):
                unit, consumed, fmt = "time", re.match(r"^\s?(?:min(?:utes?|s)?)\b", after).end(), "min"
            elif re.match(r"^\s?(?:h|hrs?|hours?)\b", after):
                consumed = re.match(r"^\s?(?:h|hrs?|hours?)\b", after).end()
                value, unit, fmt = value * 60, "time", "h"
            elif re.match(r"^\s?(?:s|secs?|seconds?)\b", after):
                consumed = re.match(r"^\s?(?:s|secs?|seconds?)\b", after).end()
                value, unit, fmt = value / 60, "time", "s"
            else:
                for pattern, name in ((r"days?|d\b", "days"), (r"weeks?|wks?", "weeks"), (r"months?|mos?\b", "months"),
                                      (r"years?|yrs?", "years"), (r"[x×](?![a-z])", "times")):
                    m = re.match(rf"^\s?(?:{pattern})", after, re.I)
                    if m:
                        unit, consumed, fmt = name, m.end(), name
                        break
            if not unit:
                noun = re.match(r"^[\s-]?([a-z][a-z-]{2,})", after)
                if noun and noun.group(1) not in STOP and noun.group(1) not in MONTHS:
                    unit, fmt = "n:" + _stem(noun.group(1).split("-")[-1]), "grouped" if "," in raw else "digits"
                elif item.col_unit:
                    unit_name = item.col_unit.replace("*60", "")
                    value = value * 60 if item.col_unit.endswith("*60") else value
                    unit, fmt = unit_name, "header"
                else:
                    unit, fmt = "plain", "grouped" if "," in raw else "digits"
        if unit == "plain":
            if 1900 <= value <= 2100 and "," not in raw and "." not in raw:
                continue  # a year
            if item.source not in {"table", "chart"} and value < 1000:
                continue  # a bare small number in prose is a count, a score or a step - too ambiguous to compare
        figure = Figure(page=item.page, item=item, raw=text[prefix_start:end + consumed].strip(), value=value, unit=unit, fmt=fmt,
                        decimals=decimals, scale=scale, start=prefix_start, end=end + consumed)
        figure.approx = bool(APPROX_BEFORE.search(text[:prefix_start]))
        figure.bound = bool(BOUND_BEFORE.search(text[:prefix_start].replace("~", "").rstrip()))
        figure.delta = bool(re.search(r"[−+±]\s?$|(?<![\w\s])-\s?$|\(\s?[−+-]\s?$", text[:prefix_start]))
        figures.append(figure)
        taken = end + consumed
    # ranges (15-49 min) and transitions (10.4 -> 25.5 min): the first number takes the second's unit
    for first, second in zip(figures, figures[1:]):
        between = first.item.text[first.end:second.start]
        if re.fullmatch(r"\s?[–—-]\s?|\s?to\s?", between) and first.unit in {"plain", ""}:
            first.unit, first.fmt, first.ranged, second.ranged = second.unit, second.fmt, True, True
            if second.unit == "time" and second.fmt == "h":
                first.value *= 60
        elif re.fullmatch(r"\s?[–—-]\s?", between):
            first.ranged = second.ranged = True
        elif re.fullmatch(r"\s?(?:→|->|⟶|to)\s?", between) and first.unit == "plain" and second.unit != "plain":
            first.unit, first.fmt = second.unit, second.fmt
            if second.unit == "time" and second.fmt == "h":
                first.value *= 60
    return figures


def _stem(word: str) -> str:
    word = word.lower().strip("-'’")
    for suffix in ("ations", "ation", "ions", "ion", "ers", "er", "ors", "or", "ings", "ing", "ies", "es", "is", "ed", "al", "s", "e", "y"):
        if len(word) - len(suffix) >= 4 and word.endswith(suffix):
            return word[: -len(suffix)]
    return word


ENTITY_RE = re.compile(r"\b([a-z][a-z-]*[a-z])(\d{1,2})\b", re.I)


def _tokens(text: str, drop: set[str] | None = None) -> list[str]:
    words = re.findall(r"[A-Za-z][A-Za-z0-9'’/-]*[A-Za-z0-9]|[A-Za-z]", text)
    out = []
    for word in words:
        low = word.lower().strip("'’")
        if low in STOP or low in MONTHS or len(low) < 2 or re.fullmatch(r"\d+", low):
            continue
        token = low if ENTITY_RE.fullmatch(low) else _stem(low)
        if drop and token in drop:
            continue
        out.append(token)
    return out


def _clause(text: str, start: int, end: int) -> tuple[str, str]:
    """The words of the figure's clause before and after it (clauses end at ; : . ! ? | · • newline and em dashes). A colon
    ends a clause after the figure but not before it: `Fixed fee: HKD 7,507,000` is labelled by what precedes the colon."""
    stops = re.compile(r"[;:!?|·•\n]|\.(?=\s|$)|\s[—–]\s|\(|\)")
    left_stops = re.compile(r"[;!?|·•\n]|\.(?=\s|$)|\s[—–]\s|\(|\)")
    left_cut = max((m.end() for m in left_stops.finditer(text, 0, start)), default=0)
    right = stops.search(text, end)
    return text[left_cut:start], text[end: right.start() if right else len(text)]


def assign_keys(pages: list[Page], figures: list[Figure]) -> None:
    """Give every figure its context key: row and column headers for a table cell or chart value; otherwise the nearest
    content words of its clause (before it when there are any, else after it). Acronyms are expanded from the deck
    (MTTR -> mean time to resolve) and words that appear on most pages carry no meaning as a key."""
    page_count = max(1, len(pages))
    document = " ".join(item.text for page in pages for item in page.items if not item.chrome)
    acronyms = _acronyms(document)
    frequency: Counter = Counter()
    for page in pages:
        frequency.update({token for item in page.items if not item.chrome for token in _tokens(item.text)})
    common = {token for token, count in frequency.items() if page_count >= 4 and count > page_count * 0.6 and not ENTITY_RE.fullmatch(token)}

    def tokens(text: str, drop_common: bool = True) -> list[str]:
        out = []
        for token in _tokens(text, common if drop_common else None):
            out.extend(acronyms.get(token, [token]))
        return out

    families = {page.number: _families(page) for page in pages}
    by_item: dict[int, list[Figure]] = defaultdict(list)
    for figure in figures:
        by_item[id(figure.item)].append(figure)
    page_of = {page.number: page for page in pages}
    for figure in figures:
        item = figure.item
        if item.source == "text" and not item.row and not item.col and _value_cell(item, by_item[id(item)]):
            row, col, caption = _grid_context(page_of[item.page], item)  # a drawn table or a KPI tile: its labels are other texts
            if row or col or caption:
                figure.rowkey = frozenset(tokens(row, False)) if row else frozenset()
                figure.colkey = frozenset(tokens(col + " " + caption, False))
                before, after = _clause(item.text, figure.start, figure.end)
                figure.key = frozenset(figure.rowkey | figure.colkey | set(tokens(re.sub(NUMBER_RE, " ", before + " " + after))))
                figure.label = " / ".join(part for part in (row, col, caption) if part)[:120]
                figure.anchors = frozenset(t for t in figure.key if ENTITY_RE.fullmatch(t))
                continue
        if item.source in {"table", "chart"} and (item.row or item.col):
            figure.rowkey, figure.colkey = frozenset(tokens(item.row, False)), frozenset(tokens(item.col, False))
            words = list(figure.rowkey | figure.colkey)  # a header is an explicit label: its words count even when common
            before, after = _clause(item.text, figure.start, figure.end)
            local = tokens(re.sub(NUMBER_RE, " ", before + " " + after))
            figure.key = frozenset(words + local)
            figure.label = " / ".join(part for part in (item.row, item.col) if part)
        else:
            before, after = _clause(item.text, figure.start, figure.end)
            siblings = [f for f in by_item[id(item)] if f is not figure]
            for other in siblings:  # a clause holding several figures: stop at the neighbour
                if other.end <= figure.start and item.text[other.end:figure.start] in before:
                    before = item.text[other.end:figure.start] if other.end >= figure.start - len(before) else before
                if other.start >= figure.end and item.text[figure.end:other.start] in after:
                    after = item.text[figure.end:other.start] if other.start <= figure.end + len(after) else after
            left = tokens(before)[-5:]
            right = tokens(after)[:5]
            words = left if len(left) >= 1 and not (len(left) == 1 and right and len(right) >= 3) else right
            trailing = right[:1] if words is left and re.match(r"^\s?[A-Za-z]", after) else []  # `5% target`, `$5.49 whole`
            figure.key = frozenset(words + trailing)
            figure.label = " ".join((before if words is left else after).split()[-8:] if words is left else after.split()[:8])
        figure.anchors = frozenset(t for t in figure.key if ENTITY_RE.fullmatch(t))
    for page in pages:  # "$5.24 -> $9.14" on a page about kirkland1 and kirkland2: before belongs to 1, after to 2
        pairs = [f for f in families[page.number] if len(f) == 2]
        if len(pairs) != 1:
            continue
        first_name, second_name = pairs[0]
        page_figures = [f for f in figures if f.page == page.number]
        for a, b in zip(page_figures, page_figures[1:]):
            if a.item is b.item and a.unit == b.unit and re.fullmatch(r"\s?(?:→|->|⟶)\s?", a.item.text[a.end:b.start]):
                if not a.anchors and not b.anchors:
                    a.key, a.anchors = a.key | {first_name}, frozenset({first_name})
                    b.key, b.anchors = b.key | {second_name}, frozenset({second_name})


def _value_cell(item: Item, figures: list[Figure]) -> bool:
    """A text that is mostly its figure (`4 h 30 min`, `2 h 25 min (-50%)`, `USD 98,920`): its label is elsewhere."""
    rest = item.text
    for figure in sorted(figures, key=lambda f: -f.start):
        rest = rest[:figure.start] + " " + rest[figure.end:]
    return bool(figures) and len(re.findall(r"[A-Za-z]{2,}", rest)) <= 4 and len(item.text) <= 60


def _label_like(candidate: Item, item: Item) -> bool:
    text = candidate.text.strip()
    return (candidate is not item and not candidate.chrome and candidate.source == "text" and 3 <= len(text) <= 80
            and len(re.findall(r"[A-Za-z]{3,}", text)) >= 1 and not re.match(r"^[~≈<>≤≥$€£¥+−-]?\s*\d", text))


def _grid_context(page: Page, item: Item) -> tuple[str, str, str]:
    """For a figure standing alone in its text: the row label to its left (the leftmost short label on its line), the column
    header above it (the nearest label aligned with it, within 320 px), and a caption directly below it (a KPI tile)."""
    top, bottom = item.y, item.y + item.h
    row = col = caption = ""
    same_row = [c for c in page.items if _label_like(c, item) and c.x + min(c.w, 420) <= item.x + 4 and item.x - c.x <= 760
                and c.h <= 3.5 * max(item.h, 1) and min(bottom, c.y + c.h) - max(top, c.y) >= 0.3 * min(item.h, c.h)]
    if same_row:
        row = min(same_row, key=lambda c: c.x).text
    above = [c for c in page.items if _label_like(c, item) and c.y + c.h <= top + 2 and top - (c.y + c.h) <= 320
             and (abs(c.ax - item.ax) <= 16 and c.anchor == item.anchor
                  or min(item.x + item.w, c.x + c.w) - max(item.x, c.x) >= 0.5 * min(item.w, c.w))]
    if above:
        nearest = max(above, key=lambda c: c.y)
        if not same_row and nearest.y + nearest.h >= top - 1.5 * item.size - 24:
            caption = nearest.text  # a label just over a KPI
        headers = [c for c in above if c.text.isupper() or str(c.weight) in {"bold", "600", "700", "800", "900"}]
        col = max(headers, key=lambda c: c.y).text if headers and same_row else ""
    if not same_row and not caption:
        below = [c for c in page.items if _label_like(c, item) and c.y >= bottom - 2 and c.y - bottom <= 1.5 * item.size + 24
                 and min(item.x + item.w, c.x + c.w) - max(item.x, c.x) >= 0.3 * min(item.w, c.w)]
        if below:
            caption = min(below, key=lambda c: c.y).text
    return row, col, caption


def _families(page: Page) -> list[tuple[str, ...]]:
    """Numbered names on a page grouped by stem: kirkland1, kirkland2 -> ("kirkland1", "kirkland2")."""
    members: dict[str, set[int]] = defaultdict(set)
    for item in page.items:
        for stem, number in ENTITY_RE.findall(item.text):
            members[stem.lower()].add(int(number))
    return [tuple(f"{stem}{n}" for n in sorted(numbers)) for stem, numbers in members.items() if len(numbers) >= 2]


def _acronyms(text: str) -> dict[str, list[str]]:
    """Acronyms that the deck spells out somewhere (MTTR <- Mean time to resolve), mapped to the phrase's key tokens."""
    found: dict[str, list[str]] = {}
    words = re.findall(r"[A-Za-z][A-Za-z-]*", text)
    initials = [w[0].upper() for w in words]
    joined = "".join(initials)
    for acronym in set(re.findall(r"\b[A-Z]{2,6}\b", text)):
        for match in re.finditer(acronym, joined):
            phrase = words[match.start(): match.start() + len(acronym)]
            if phrase and any(w[0].isupper() or w[0].islower() for w in phrase) and not all(w.isupper() for w in phrase):
                tokens = [_stem(w) for w in phrase if w.lower() not in STOP]
                if len(tokens) >= 2:
                    found[acronym.lower()] = tokens
                    found[_stem(acronym.lower())] = tokens
                    break
    return found


def _compatible(a: Figure, b: Figure) -> bool:
    """Equal, or equal within the rounding the less precise one was written with (7.5m ~ 7,507,000; ~28 ~ 27.6)."""
    if math.isclose(a.value, b.value, rel_tol=1e-9, abs_tol=1e-9):
        return True
    coarse = a if (a.decimals - math.log10(a.scale)) < (b.decimals - math.log10(b.scale)) else b
    step = coarse.scale * 10 ** (-coarse.decimals)
    if abs(a.value - b.value) <= step / 2 + 1e-9:
        return True
    if a.approx or b.approx:
        return abs(a.value - b.value) <= 0.1 * max(abs(a.value), abs(b.value))
    return False


def _keys_match(a: Figure, b: Figure) -> bool:
    """Do two figures label the same quantity?
    - named runs or items (kirkland1): the same names, and the rest of the labels agree or one has none;
    - two table cells: the same row (a row header names an entity: `Delivery Manager` is not `Manager`) and a shared column word;
    - a table cell and prose: the prose names most of the row;
    - two prose figures: at least two words in common, and most of the shorter label."""
    if a.anchors or b.anchors:
        if a.anchors != b.anchors:
            return False
        rest_a, rest_b = a.key - a.anchors, b.key - b.anchors
        return not rest_a or not rest_b or bool(rest_a & rest_b)
    if a.rowkey and b.rowkey:
        if a.rowkey != b.rowkey and len(a.rowkey & b.rowkey) / len(a.rowkey | b.rowkey) < 0.75:
            return False
        return not a.colkey or not b.colkey or bool(a.colkey & b.colkey)
    if a.rowkey or b.rowkey:
        row, other = (a.rowkey, b.key) if a.rowkey else (b.rowkey, a.key)
        shared = row & other
        return len(shared) >= 2 and len(shared) >= 0.66 * len(row)
    shared = a.key & b.key
    if len(shared) < 2:
        return False
    return len(shared) / min(len(a.key), len(b.key)) >= 0.75


def _digits(figure: Figure) -> str:
    number = NUMBER_RE.search(figure.raw)
    return re.sub(r"\D", "", number.group(0)) if number else ""


def _near_miss(a: Figure, b: Figure) -> bool:
    """Two figures that are the same long number but for one digit (98,920 / 98,720): a copy error, whatever their labels -
    independent quantities rarely look alike. Needs four or more digits, the same unit and at most 5% apart."""
    if a.unit != b.unit or a.unit in {"%", "pp", "time"} or a.scale != b.scale:
        return False
    da, db = _digits(a), _digits(b)
    if len(da) != len(db) or len(da.lstrip("0")) < 4 or da == db:
        return False
    if sum(x != y for x, y in zip(da, db)) != 1 or abs(a.value - b.value) > 0.05 * max(abs(a.value), abs(b.value)):
        return False
    if a.item.source != "text" and b.item.source != "text":
        return False  # two table cells a digit apart are usually two line items
    if a.anchors != b.anchors or (a.rowkey and b.rowkey and not (a.rowkey & b.rowkey)):
        return False
    return True


def _comparable(f: Figure) -> bool:
    return not (f.bound or f.delta or f.ranged) and f.unit not in {"", "times"} and bool(f.key)


def figure_conflicts(pages: list[Page], figures: list[Figure], spec_text: str = "") -> list[dict]:
    """Pages that give one quantity values with nothing in common. A page that states both values (a baseline and its
    target) is detail, not conflict; only figures within 50% of each other are compared, so a before/after pair on two
    pages whose keys happen to match is not reported."""
    usable = [f for f in figures if _comparable(f)]
    by_page: dict[int, list[Figure]] = defaultdict(list)
    for figure in usable:
        by_page[figure.page].append(figure)
    findings: dict[tuple, dict] = {}
    numbers = sorted(by_page)
    near = [f for f in figures if not (f.bound or f.delta or f.ranged) and f.unit not in {"", "times"}]
    near_page: dict[int, list[Figure]] = defaultdict(list)
    for figure in near:
        near_page[figure.page].append(figure)
    for i, pa in enumerate(numbers):
        for pb in numbers[i + 1:]:
            for fa in by_page[pa]:
                matches = [fb for fb in by_page[pb] if fb.unit == fa.unit and _keys_match(fa, fb)]
                if not matches or any(_compatible(fa, fb) for fb in matches):
                    continue
                fb = min(matches, key=lambda f: abs(f.value - fa.value))
                if abs(fa.value - fb.value) > 0.5 * max(abs(fa.value), abs(fb.value)):
                    continue
                back = [f for f in by_page[pa] if f.unit == fb.unit and _keys_match(f, fb)]
                if any(_compatible(f, fb) for f in back):
                    continue  # page A also states B's value in the same context: two states of one thing
                own_a, own_b = fa.key - fb.key - fa.anchors, fb.key - fa.key - fb.anchors
                agree = any(_compatible(x, y) for x in back + [fa] for y in matches)
                if own_a and own_b and agree:
                    continue  # the pages agree on one state (12% today) and label the others differently (7% year one / 5% target)
                key = (pa, pb, round(fa.value, 6), round(fb.value, 6), fa.unit)
                if key in findings:
                    continue
                findings[key] = _conflict_finding(fa, fb, spec_text)
                if agree and (own_a or own_b):  # one side names a scope the other leaves out: probably two states, say so
                    findings[key]["severity"] = FLAGGED
                    findings[key]["message"] += " The pages agree on another value of it; one page names a scope the other does not - probably two states, unlabelled on one page."
    pages_near = sorted(near_page)
    for i, pa in enumerate(pages_near):  # one digit apart, whatever the labels say
        for pb in pages_near[i + 1:]:
            for fa in near_page[pa]:
                for fb in near_page[pb]:
                    if not _near_miss(fa, fb):
                        continue
                    if any(_compatible(f, fb) and f.unit == fb.unit for f in near_page[pa]) or any(_compatible(f, fa) and f.unit == fa.unit for f in near_page[pb]):
                        continue  # each page also states the other's value: two quantities
                    key = (pa, pb, round(fa.value, 6), round(fb.value, 6), fa.unit)
                    if key not in findings:
                        findings[key] = _conflict_finding(fa, fb, spec_text, near=True)
    return list(findings.values())


def _in_spec(figure: Figure, spec_text: str) -> bool:
    """The plan states this figure (only a distinctive one: a bare small number is everywhere in a spec)."""
    if not spec_text or (figure.unit == "plain" and figure.value < 1000) or len(re.sub(r"\D", "", figure.raw)) < 3:
        return False
    candidates = {figure.raw, figure.item.text[figure.start:figure.end]}
    number = re.search(NUMBER_RE, figure.raw)
    if number:
        candidates.add(number.group(0))
    return any(c and re.search(r"(?<![\d.,])" + re.escape(c) + r"(?![\d])", spec_text) for c in candidates if len(c) >= 2)


def _conflict_finding(fa: Figure, fb: Figure, spec_text: str, near: bool = False) -> dict:
    in_a, in_b = _in_spec(fa, spec_text), _in_spec(fb, spec_text)
    context = " ".join(sorted(fa.anchors)) or (fa.label or fb.label)
    if in_a and not in_b:
        repair = {fb.page: f"P{fb.page:02d} says {fb.raw}; the plan (design_spec.md) and P{fa.page:02d} say {fa.raw} - use {fa.raw}, or state why this page's scope differs."}
        direction = f"P{fb.page:02d} deviates from the plan's figure {fa.raw}"
    elif in_b and not in_a:
        repair = {fa.page: f"P{fa.page:02d} says {fa.raw}; the plan (design_spec.md) and P{fb.page:02d} say {fb.raw} - use {fb.raw}, or state why this page's scope differs."}
        direction = f"P{fa.page:02d} deviates from the plan's figure {fb.raw}"
    else:
        text = (f"P{fa.page:02d} says {fa.raw} and P{fb.page:02d} says {fb.raw} for {context!s}: reconcile both pages to the source - "
                "state the scope on each, or use one figure.")
        repair = {fa.page: text, fb.page: text}
        direction = "state the scope on each page or use one figure (reconcile to the source)"
    return {"kind": "FIGURE_CONFLICT", "severity": CERTAIN, "pages": sorted({fa.page, fb.page}),
            "elements": [_element(fa.item, fa.raw), _element(fb.item, fb.raw)],
            "message": f"P{fa.page:02d} says {fa.raw} ({_label(fa)}); P{fb.page:02d} says {fb.raw} ({_label(fb)}) - "
                       + ("one digit apart and neither page states the other: a copy error or two scopes." if near else "the same quantity, two values."),
            "fix": f"P{fa.page:02d} says {fa.raw} and P{fb.page:02d} says {fb.raw} for {context}: {direction}.",
            "repair": {str(k): v for k, v in repair.items()}}


def _label(figure: Figure) -> str:
    parts = sorted(figure.anchors) + ([figure.label] if re.search(r"[A-Za-z]{2}", figure.label) else [])
    return "; ".join(p for p in parts if p)[:90] or "no label"


def _element(item: Item, focus: str = "") -> dict:
    text = item.text.replace("\n", " ")
    at = text.find(focus) if focus else -1
    excerpt = text[max(0, at - 60): at + len(focus) + 60] if at >= 0 else text[:140]
    return {"page": item.page, "element": item.eid, "excerpt": excerpt.strip()}


FORMAT_UNITS = {"USD", "HKD", "SGD", "AUD", "CAD", "NZD", "EUR", "GBP", "CHF", "JPY", "CNY", "INR", "time", "plain", "%"}


def format_drift(figures: list[Figure]) -> list[dict]:
    """One value written two ways (7.507m / 7,507,000; 4 h 50 min / 290 min; USD / $) on the deck's pages."""
    groups: dict[tuple, list[Figure]] = defaultdict(list)
    for figure in figures:
        if figure.ranged or figure.unit not in FORMAT_UNITS or figure.fmt == "header":
            continue  # a table column that states its unit in the header writes bare numbers by convention
        distinctive = (figure.unit == "time" and figure.value >= 60) or (figure.unit not in {"time", "%"} and abs(figure.value) >= 1000)
        if not distinctive:
            continue
        groups[(figure.unit, round(figure.value, 4))].append(figure)
    findings = []
    for (unit, value), members in groups.items():
        forms: dict[str, list[Figure]] = defaultdict(list)
        for figure in members:
            forms[_form(figure)].append(figure)
        pages = {f.page for f in members}
        if len(forms) < 2 or len(pages) < 2:
            continue
        ranked = sorted(forms.items(), key=lambda kv: (-len({f.page for f in kv[1]}), -len(kv[1])))
        keep = ranked[0][1][0].raw
        described = "; ".join(f"{fs[0].raw!r} on " + ", ".join(f"P{p:02d}" for p in sorted({f.page for f in fs})) for _, fs in ranked)
        others = sorted({f.page for _, fs in ranked[1:] for f in fs})
        findings.append({"kind": "FORMAT_DRIFT", "severity": FLAGGED, "pages": sorted(pages),
                         "elements": [_element(fs[0].item, fs[0].raw) for _, fs in ranked],
                         "message": f"one value, {len(forms)} written forms: {described}.",
                         "fix": f"write it one way on every page - {keep!r} (the most used form) - on " + ", ".join(f"P{p:02d}" for p in others) + ".",
                         "repair": {str(p): f"Write {keep!r} in the deck's one form (this page writes the same value as "
                                    + ", ".join(repr(f.raw) for f in members if f.page == p) + ")." for p in others}})
    return findings


def _form(figure: Figure) -> str:
    if figure.unit == "time":
        return figure.fmt
    return figure.fmt.replace("+digits", "+grouped") if figure.value < 10000 else figure.fmt


# --------------------------------------------------------------------------------------------------------------------------
# terms
# --------------------------------------------------------------------------------------------------------------------------

TERM_RE = re.compile(r"\b([A-Z][a-z]+(?:[- ](?:[A-Z][a-z]+|[a-z]{3,})){1,2})\b")
GENERIC_HEADS = {"the", "a", "an", "this", "that", "our", "your", "their", "each", "every", "all", "no", "one", "two", "three", "if", "when",
                 "where", "why", "how", "what", "who", "not", "and", "or", "but", "for", "from", "with", "without", "to", "in", "on", "by",
                 "at", "as", "it", "its", "is", "are", "we", "you", "they", "then", "so", "use", "run", "see", "add", "set", "keep", "make"}


def glossary_terms(spec_text: str) -> list[str]:
    """`## Glossary` or `## Canonical terms` in design_spec.md: `- Term: definition` (or `- **Term**: definition`) lines."""
    match = re.search(r"^#{1,4}\s*(?:[IVX]+\.\s*)?(?:Glossary|Canonical terms?)\b[^\n]*\n(.*?)(?=^#{1,4}\s|\Z)", spec_text or "", re.M | re.S | re.I)
    if not match:
        return []
    terms = []
    for line in match.group(1).splitlines():
        m = re.match(r"^\s*[-*]\s+\**([^:*]+?)\**\s*:", line)
        if m and 1 <= len(m.group(1).split()) <= 6:
            terms.append(m.group(1).strip())
    return terms


def _term_key(term: str) -> tuple[str, ...]:
    return tuple(_stem(w) for w in re.split(r"[\s-]+", term.lower()) if w)


def term_drift(pages: list[Page], spec_text: str = "") -> list[dict]:
    """Near-duplicate names on different pages: the same words up to case, hyphenation and suffix (Diagnose/Diagnosis Agent),
    found by clustering Title-Case terms - or, with a glossary, every variant of a canonical term."""
    occurrences: dict[tuple, dict[str, set[int]]] = defaultdict(lambda: defaultdict(set))
    elements: dict[tuple, dict[str, Item]] = defaultdict(dict)
    texts = [(page.number, item) for page in pages for item in page.items if not item.chrome and item.source != "chart"]
    texts += [(page.number, item) for page in pages for item in _stacked_labels(page)]  # `Diagnose` over `agent` reads as one name
    named_mid_sentence: set[tuple] = set()
    for number, item in texts:
        if item.text.isupper():
            continue
        for match in TERM_RE.finditer(item.text):
            term = re.sub(r"(?:[\s-][a-z]+)+$", "", match.group(1))  # `Review Agent checks` -> `Review Agent`
            while term.split(" ", 1)[0].lower() in GENERIC_HEADS and " " in term:
                term = term.split(" ", 1)[1]  # `The Review Agent` -> `Review Agent`
            words = re.split(r"[\s-]+", term)
            if len(words) < 2 or words[0].lower() in GENERIC_HEADS or words[-1].lower() in STOP:
                continue
            if sum(1 for w in words if w[0].isupper()) < 2 or any(w.lower() in MONTHS for w in words):
                continue  # a name is Title Case (Review Agent, Decision Board); `City-selected` is a phrase
            key = _term_key(term)
            start = match.start() + match.group(1).find(term)
            if re.search(r"[^\s.!?:;]\s*$", item.text[:start]) or re.match(r"\s*(?:$|[·|,/(:–—-])", item.text[start + len(term):]):
                named_mid_sentence.add(key)  # `the Review Agent checks`, or a label `Diagnosis Agent · ...`: Title Case that is a name
            occurrences[key][term].add(number)
            elements[key].setdefault(term, item)
    for key in [k for k in occurrences if k not in named_mid_sentence]:
        del occurrences[key]  # only ever capitalised at the start of a label or sentence (`Follow City ...`): not evidence of a name
    canonical = {_term_key(t): t for t in glossary_terms(spec_text)}
    for key in set(occurrences) | set(canonical):  # every surface form of a name, in any case: `Reviewer agent` for `Review Agent`
        pattern = re.compile(r"\b" + r"[\s-]?".join(re.escape(stem) + r"[a-z]{0,6}" for stem in key) + r"\b", re.I)
        for number, item in texts:
            for match in pattern.finditer(item.text):
                found = match.group(0)
                if _term_key(found) == key or (key in canonical and _loose_key(found) == _loose_key(canonical[key])):
                    occurrences[key][found].add(number)
                    elements[key].setdefault(found, item)
    findings = []
    for key, forms in occurrences.items():
        distinct = {}
        for form, numbers in forms.items():
            norm = re.sub(r"\s+", " ", form.lower())
            distinct.setdefault(norm, [form, set()])
            distinct[norm][1] |= numbers
        pages_all = set().union(*[v[1] for v in distinct.values()])
        target = canonical.get(key)
        if target:
            want = re.sub(r"\s+", " ", target.lower())  # case is typography (a caps label); words, hyphens and suffixes are the term
            wrong = {n: v for n, v in distinct.items() if n != want and not _plural_only([n, want])}
            if not wrong:
                continue
            listing = "; ".join(f"{v[0]!r} on " + ", ".join(f"P{p:02d}" for p in sorted(v[1])) for v in wrong.values())
            findings.append({"kind": "TERM_DRIFT", "severity": FLAGGED, "pages": sorted(set().union(*[v[1] for v in wrong.values()])),
                             "elements": [_element(elements[key][v[0]], v[0]) for v in wrong.values()],
                             "message": f"the glossary term {target!r} is written differently: {listing}.",
                             "fix": f"write {target!r} exactly, as design_spec.md's glossary does.",
                             "repair": {str(p): f"Write the glossary term {target!r} exactly (this page has " + ", ".join(repr(v[0]) for v in wrong.values() if p in v[1]) + ")."
                                        for v in wrong.values() for p in v[1]}})
            continue
        if len(distinct) < 2 or len(pages_all) < 2:
            continue
        surfaces = list(distinct)
        if _plural_only(surfaces):
            continue
        ranked = sorted(distinct.values(), key=lambda v: (-len(v[1]), not all(w[:1].isupper() for w in re.split(r"[\s-]+", v[0])), v[0]))
        keep = ranked[0][0]
        listing = "; ".join(f"{v[0]!r} on " + ", ".join(f"P{p:02d}" for p in sorted(v[1])) for v in ranked)
        others = sorted(set().union(*[v[1] for v in ranked[1:]]) - ranked[0][1]) or sorted(set().union(*[v[1] for v in ranked[1:]]))
        findings.append({"kind": "TERM_DRIFT", "severity": FLAGGED, "pages": sorted(pages_all),
                         "elements": [_element(elements[key][v[0]], v[0]) for v in ranked],
                         "message": f"one name, {len(ranked)} forms: {listing}.",
                         "fix": f"if these are one thing, name it one way ({keep!r}, the most used form) on " + ", ".join(f"P{p:02d}" for p in others) + ".",
                         "repair": {str(p): f"Name it {keep!r} as the rest of the deck does, if it is the same thing." for p in others}})
    return findings


def _stacked_labels(page: Page) -> list[Item]:
    """Short labels drawn one under the other in one group and aligned (a step name over its owner): the words a reader
    takes together. PowerPoint's text-in-shapes pass joins them into one text frame; the SVG keeps them apart."""
    out = []
    items = [i for i in page.items if i.source == "text" and not i.chrome and len(i.text.split()) <= 3 and "\n" not in i.text]
    for upper in items:
        below = [i for i in items if i is not upper and i.eid == upper.eid and abs(i.ax - upper.ax) <= 4 and i.anchor == upper.anchor
                 and 0 < i.y - upper.y <= 1.8 * max(upper.size, i.size)]
        if below:
            lower = min(below, key=lambda i: i.y)
            out.append(Item(page=page.number, eid=upper.eid, text=f"{upper.text} {lower.text}", x=upper.x, y=upper.y, w=max(upper.w, lower.w),
                            h=lower.y + lower.h - upper.y, size=upper.size, ax=upper.ax, anchor=upper.anchor))
    return out


def _loose_key(term: str) -> str:
    return re.sub(r"[\s-]+", "", term.lower())


def _plural_only(surfaces: list[str]) -> bool:
    """Forms that differ only by an inflection of the last word (plural, possessive, -ing, -ed), or by case, are grammar,
    not drift: `City-weighted scores` / `City-weighted scoring`."""
    def base(s: str) -> str:
        words = s.split()
        return " ".join(words[:-1] + [_stem(re.sub(r"(?:'s|’s)$", "", words[-1]))])
    return len({base(s) for s in surfaces}) < 2


# --------------------------------------------------------------------------------------------------------------------------
# chrome
# --------------------------------------------------------------------------------------------------------------------------


def _is_cover(page: Page, pages: list[Page]) -> bool:
    name = (page.stem + " " + page.layout).lower()
    return any(word in name for word in ("cover", "title", "closing", "section", "divider", "agenda", "end")) or page.number == min(p.number for p in pages)


def _chrome_parts(page: Page) -> list[tuple[str, tuple, dict]]:
    """(kind, slot, style) for every chrome element of a page. Slots are positions on an 8 px grid; a text slot is found by
    its band and alignment edge, so a section label that changes wording stays the same slot."""
    parts = []
    for item in page.items:
        if not item.chrome or item.layer == "layout":  # a layout's own furniture differs by layout by design
            continue
        text = item.text.strip()
        folio = bool(re.fullmatch(r"\d{1,3}", text))
        slot = _text_slot(item, page)
        parts.append(("text", slot, {"text": "#" if folio else re.sub(r"\d+", "#", text), "size": item.size, "fill": item.fill,
                                     "weight": item.weight, "folio": text if folio else "", "eid": item.eid, "band": item.band,
                                     "edge": item.y + item.h <= page.height * 0.085 or item.y >= page.height * 0.915}))
    for shape in page.shapes:
        if not shape.chrome or shape.layer == "layout":
            continue
        parts.append((shape.tag, _shape_slot(shape), {"fill": shape.fill, "stroke": shape.stroke, "eid": shape.eid, "edge": True}))
    return parts


def _shape_slot(shape: Shape) -> tuple:
    return (shape.tag, round(shape.x / 8), round(shape.y / 8), round(shape.w / 8), round(shape.h / 8))


def _drawn_anyway(page: Page, slot: tuple) -> bool:
    """The page draws the element, only outside the chrome group (a title rule kept in the title group): no visible drift."""
    if slot[0] == "text":
        return any(item.source == "text" and _text_slot(item, page) == slot for item in page.items)
    return any(_shape_slot(shape) == slot or (shape.tag == slot[0] and _near(slot, _shape_slot(shape))) for shape in page.shapes)


def _text_slot(item: Item, page: Page) -> tuple:
    band = "top" if item.y < page.height / 2 else "bottom"
    if re.fullmatch(r"\d{1,3}", item.text.strip()):
        return ("text", band, "folio", round(item.ax / 48), item.anchor)
    return ("text", band, round(item.y / 8), round(item.ax / 12), item.anchor)


def chrome_drift(pages: list[Page]) -> list[dict]:
    """Body pages whose chrome differs from the majority of their master: a missing or extra element or a folio out of
    sequence is certain; a different colour, size or wording is flagged."""
    findings: list[dict] = []
    groups: dict[tuple, list[Page]] = defaultdict(list)
    for page in pages:  # pages of one master and one ground: a dark closer or divider is its own kind of page
        if not _is_cover(page, pages):
            groups[(page.master, _ground(page))].append(page)
    for master, body in groups.items():
        if len(body) < 3:
            continue
        parts = {page.number: _chrome_parts(page) for page in body}
        if not any(parts.values()):
            continue
        slot_pages: dict[tuple, set[int]] = defaultdict(set)
        for number, entries in parts.items():
            for _, slot, _ in entries:
                slot_pages[slot].add(number)
        quorum = len(body) / 2
        common = {slot for slot, numbers in slot_pages.items() if len(numbers) > quorum}
        for page in body:
            entries = parts[page.number]
            present = {slot for _, slot, _ in entries}
            missing = [slot for slot in common if slot not in present]
            unshared = [(kind, slot, style) for kind, slot, style in entries if slot not in common]
            moved = []
            for slot in list(missing):  # the same element a little elsewhere is a style difference, not a missing one
                near = next((e for e in unshared if e[0] == slot[0] and _near(slot, e[1])), None)
                if near:
                    missing.remove(slot)
                    unshared.remove(near)
                    moved.append((slot, near))
            extra = [e for e in unshared if not e[2].get("band") and len(slot_pages[e[1]]) == 1]  # content that merely sits low or high is not chrome
            missing = [slot for slot in missing if not _drawn_anyway(page, slot)]
            for slot in missing:
                sample = _style_for(parts, slot)
                findings.append({"kind": "CHROME_DRIFT", "severity": CERTAIN, "pages": [page.number],
                                 "elements": [{"page": page.number, "element": sample.get("eid", ""), "excerpt": sample.get("text", slot[0])}],
                                 "message": f"P{page.number:02d} lacks a chrome element that {len(slot_pages[slot])} of {len(body)} body pages carry: "
                                            f"{slot[0]} {sample.get('text') or sample.get('fill') or ''!s}".strip() + ".",
                                 "fix": f"P{page.number:02d}: restore the chrome element exactly as the other body pages draw it (copy it from the template).",
                                 "repair": {str(page.number): f"Your page lacks the deck's chrome element {slot[0]} "
                                            f"{sample.get('text') or sample.get('fill') or ''!s} that the other body pages carry: copy it from the template exactly."}})
            if len(extra) > 4:  # a page that wrapped its own content in the chrome group: one finding, not one per element
                findings.append({"kind": "CHROME_DRIFT", "severity": CERTAIN, "pages": [page.number],
                                 "elements": [{"page": page.number, "element": style.get("eid", ""), "excerpt": str(style.get("text", kind))[:100]}
                                              for kind, _, style in extra[:3]],
                                 "message": f"P{page.number:02d} marks {len(extra)} elements as chrome that no other body page has - its own content sits "
                                            "inside the chrome group (it would be promoted to the slide layout).",
                                 "fix": f"P{page.number:02d}: keep only the template's chrome in the chrome group; move the page's content out of it.",
                                 "repair": {str(page.number): f"{len(extra)} of your page's own elements sit inside the chrome group (<g id='chrome'> or "
                                            "data-pptx-role='chrome'): keep only the template's header/footer there and move the content out."}})
                extra = []
            for kind, slot, style in extra:
                if kind == "text" and len(common) == 0:
                    continue
                inside = kind == "text" and not style.get("edge")  # marked as chrome but drawn where content goes: a markup slip
                findings.append({"kind": "CHROME_DRIFT", "severity": FLAGGED if inside else CERTAIN, "pages": [page.number],
                                 "elements": [{"page": page.number, "element": style.get("eid", ""), "excerpt": style.get("text", kind)}],
                                 "message": f"P{page.number:02d} carries a chrome element no other body page has: {kind} {style.get('text') or style.get('fill') or ''!s}".strip() + ".",
                                 "fix": f"P{page.number:02d}: remove it, or move it out of the header/footer bands if it is page content.",
                                 "repair": {str(page.number): f"Your page has a header/footer element no other page has ({kind} "
                                            f"{style.get('text') or style.get('fill') or ''!s}): remove it or move it into the body."}})
            for slot, (kind, _, style) in moved:
                findings.append({"kind": "CHROME_DRIFT", "severity": FLAGGED, "pages": [page.number],
                                 "elements": [{"page": page.number, "element": style.get("eid", ""), "excerpt": style.get("text", kind)}],
                                 "message": f"P{page.number:02d}'s chrome {kind} sits away from where the other body pages put it.",
                                 "fix": f"P{page.number:02d}: align it with the template position.",
                                 "repair": {str(page.number): "A header/footer element is out of its template position: align it with the template."}})
            for kind, slot, style in entries:  # style of shared slots
                if slot not in common:
                    continue
                peers = [s for n, es in parts.items() if n != page.number for k, sl, s in es if sl == slot]
                for attribute in ("fill", "size", "weight", "stroke", "text"):
                    values = Counter(str(p.get(attribute, "")) for p in peers)
                    if not values:
                        continue
                    usual, count = values.most_common(1)[0]
                    mine = str(style.get(attribute, ""))
                    if attribute == "text" and (values.get(mine) or count < 0.75 * len(peers)):
                        continue  # a section tracker or source line: its wording is the page's own (or its section's)
                    if mine != usual and count >= max(2, 0.6 * len(peers)):
                        findings.append({"kind": "CHROME_DRIFT", "severity": FLAGGED, "pages": [page.number],
                                         "elements": [{"page": page.number, "element": style.get("eid", ""), "excerpt": str(style.get("text", kind))[:120]}],
                                         "message": f"P{page.number:02d}'s chrome {kind} has {attribute} {mine!r} where {count} other body pages have {usual!r}.",
                                         "fix": f"P{page.number:02d}: use {attribute} {usual!r} as the other body pages do.",
                                         "repair": {str(page.number): f"A header/footer {kind} has {attribute} {mine!r}; the other pages use {usual!r} - match them."}})
        findings.extend(_folio_findings(body, parts))
        findings.extend(_partial_furniture(body, common))
    return findings


def _partial_furniture(body: list[Page], common: set) -> list[dict]:
    """The same words at the same place near the top or bottom edge on some body pages but not most (a data-class stamp on
    two pages of eight): chrome that drifted in or out. A slot every page has (a section tracker) is not this."""
    slots: dict[tuple, dict[int, Item]] = defaultdict(dict)
    for page in body:
        for item in page.items:
            if item.source != "text" or item.layer == "layout" or len(item.text) < 6 or not re.search(r"[A-Za-z]{3}", item.text):
                continue
            if not (item.y + item.h <= page.height * 0.2 or item.y >= page.height * 0.8):
                continue
            if item.chrome and _text_slot(item, page) in common:
                continue
            slots[(re.sub(r"\d+", "#", " ".join(item.text.split()).lower()), round(item.ax / 16), item.anchor)][page.number] = item
    findings = []
    for (text, _, _), where in slots.items():
        if max(i.y for i in where.values()) - min(i.y for i in where.values()) > 20:
            continue
        if 2 <= len(where) <= len(body) / 2 - 0.5:
            pages = sorted(where)
            missing = [p.number for p in body if p.number not in where]
            sample = where[pages[0]]
            findings.append({"kind": "CHROME_DRIFT", "severity": FLAGGED, "pages": pages,
                             "elements": [_element(where[n]) for n in pages[:4]],
                             "message": f"{sample.text[:80]!r} sits at the same edge position on " + ", ".join(f"P{n:02d}" for n in pages)
                                        + f" but not on the other {len(missing)} body pages - chrome on some pages only.",
                             "fix": "put it on every body page (in the template's chrome) or on none: " + ", ".join(f"P{n:02d}" for n in pages) + ".",
                             "repair": {str(n): f"{sample.text[:80]!r} appears at this edge position on only some body pages: drop it unless every body page "
                                        "carries it (then it belongs in the template's chrome)." for n in pages}})
    return findings


def _ground(page: Page) -> str:
    """The fill of the rectangle that covers the page (a full-bleed ground), if any."""
    full = [shape for shape in page.shapes if shape.tag == "rect" and shape.w * shape.h >= 0.9 * page.width * page.height
            and shape.fill and shape.fill not in {"none", "transparent"}]
    return _colour(full[-1].fill) if full else ""


def _near(a: tuple, b: tuple) -> bool:
    if a[0] == "text" and b[0] == "text":
        if a[1] != b[1] or (a[2] == "folio") != (b[2] == "folio"):
            return False
        return abs(a[3] - b[3]) <= 2 if a[2] == "folio" else abs(a[2] - b[2]) <= 4 and abs(a[3] - b[3]) <= 8
    if a[0] != b[0]:
        return False
    if a[4] <= 1 and b[4] <= 1 and abs(a[1] - b[1]) <= 3 and abs(a[3] - b[3]) <= 3:
        return abs(a[2] - b[2]) <= 8  # a rule at another height (a taller title zone): moved, not missing
    return all(abs(x - y) <= 3 for x, y in zip(a[1:], b[1:]) if isinstance(x, (int, float)))


def _style_for(parts: dict, slot: tuple) -> dict:
    for entries in parts.values():
        for _, sl, style in entries:
            if sl == slot:
                return style
    return {}


def _folio_findings(body: list[Page], parts: dict) -> list[dict]:
    folios = {}
    for page in body:
        values = [style["folio"] for _, slot, style in parts[page.number] if style.get("folio")]
        if len(values) == 1:
            folios[page.number] = int(values[0])
    if len(folios) < 3:
        return []
    offset, count = Counter(folio - number for number, folio in folios.items()).most_common(1)[0]
    if count < len(folios) / 2:
        return []
    findings = []
    for number, folio in sorted(folios.items()):
        if folio - number != offset:
            expected = number + offset
            findings.append({"kind": "CHROME_DRIFT", "severity": CERTAIN, "pages": [number],
                             "elements": [{"page": number, "element": "folio", "excerpt": str(folio)}],
                             "message": f"P{number:02d}'s folio says {folio}; the sequence says {expected}.",
                             "fix": f"P{number:02d}: set the folio to {expected}.",
                             "repair": {str(number): f"Your page's folio (page number) says {folio}; it must say {expected}."}})
    return findings


# --------------------------------------------------------------------------------------------------------------------------
# markers
# --------------------------------------------------------------------------------------------------------------------------


def marker_drift(pages: list[Page]) -> list[dict]:
    """A short label (Gate 2, Decision) keyed by a small filled shape just left of it or under it, drawn with different
    fills on different pages."""
    seen: dict[str, dict[str, set[int]]] = defaultdict(lambda: defaultdict(set))
    where: dict[tuple, Item] = {}
    for page in pages:
        markers = [s for s in page.shapes if not s.chrome and 3 <= s.w <= 28 and 3 <= s.h <= 28 and s.fill
                   and s.fill not in {"none", "#fff", "#ffffff", "white", "transparent"} and not s.fill.startswith("url(")]
        for item in page.items:
            if item.chrome or item.source != "text":
                continue
            label = re.split(r"\s[·—–|]\s|:\s", " ".join(item.text.split()))[0].strip()
            if not (2 <= len(label) <= 32 and re.search(r"[A-Za-z]", label) and len(label.split()) <= 5):
                continue
            mid = item.y + min(item.h, item.size * 1.2) / 2
            best = None
            for shape in markers:
                cy = shape.y + shape.h / 2
                if abs(cy - mid) > max(shape.h, item.size) * 0.8:
                    continue
                gap = item.x - (shape.x + shape.w)
                if -2 <= gap <= max(18.0, 1.2 * item.size):
                    if best is None or gap < best[0]:
                        best = (gap, shape)
            if best:
                norm = re.sub(r"\s+", " ", label.lower())
                seen[norm][_colour(best[1].fill)].add(page.number)
                where[(norm, page.number)] = item
    findings = []
    for label, fills in seen.items():
        pages_all = set().union(*fills.values())
        if len(fills) < 2 or len(pages_all) < 2:
            continue
        ranked = sorted(fills.items(), key=lambda kv: -len(kv[1]))
        listing = "; ".join(f"{fill} on " + ", ".join(f"P{p:02d}" for p in sorted(ps)) for fill, ps in ranked)
        keep = ranked[0][0] if len(ranked[0][1]) > len(ranked[1][1]) else "one fill (the deck's legend colour)"
        others = sorted(set().union(*[ps for _, ps in ranked[1:]])) if len(ranked[0][1]) > len(ranked[1][1]) else sorted(pages_all)
        findings.append({"kind": "MARKER_DRIFT", "severity": FLAGGED, "pages": sorted(pages_all),
                         "elements": [_element(where[(label, p)]) for p in sorted(pages_all)][:4],
                         "message": f"the marker beside {label!r} has different fills: {listing}.",
                         "fix": f"draw the {label!r} marker in {keep} on " + ", ".join(f"P{p:02d}" for p in others) + " as the other pages do.",
                         "repair": {str(p): f"The marker beside {label!r} must use {keep} as on the other pages." for p in others}})
    return findings


def _colour(fill: str) -> str:
    fill = fill.strip().lower()
    if re.fullmatch(r"#[0-9a-f]{3}", fill):
        fill = "#" + "".join(c * 2 for c in fill[1:])
    return fill


# --------------------------------------------------------------------------------------------------------------------------
# the check
# --------------------------------------------------------------------------------------------------------------------------


def check(target: Path, spec: Path | None = None) -> dict:
    pages, kind = load_pages(target)
    spec_text = spec.read_text(encoding="utf-8", errors="replace") if spec and spec.is_file() else ""
    figures = [figure for page in pages for item in page.items if not item.chrome for figure in extract_figures(item)]
    assign_keys(pages, figures)
    findings = figure_conflicts(pages, figures, spec_text) + format_drift(figures) + term_drift(pages, spec_text)
    if kind == "svg":
        findings += chrome_drift(pages) + marker_drift(pages)
    order = {"FIGURE_CONFLICT": 0, "CHROME_DRIFT": 1, "FORMAT_DRIFT": 2, "TERM_DRIFT": 3, "MARKER_DRIFT": 4}
    findings.sort(key=lambda f: (f["severity"] != CERTAIN, order.get(f["kind"], 9), f["pages"]))
    for number, finding in enumerate(findings, start=1):
        finding["id"] = f"C{number:02d}"
    counts = Counter(f["kind"] for f in findings)
    return {"target": str(target), "input": kind, "pages": [{"number": p.number, "stem": p.stem} for p in pages],
            "glossary": glossary_terms(spec_text), "figures": len(figures),
            "certain": sum(1 for f in findings if f["severity"] == CERTAIN), "flagged": sum(1 for f in findings if f["severity"] == FLAGGED),
            "counts": dict(counts), "findings": findings}


def render_markdown(report: dict, title: str = "") -> str:
    lines = [f"# Deck consistency{' - ' + title if title else ''}", "",
             f"Deterministic cross-page check over {len(report['pages'])} page(s) ({report['input']}); {report['figures']} figures read. "
             f"**{report['certain']} certain**, {report['flagged']} flagged.", "",
             "CERTAIN items are measured facts: one quantity with two values, a missing or extra chrome element, a folio out of sequence. "
             "FLAGGED items are probably drift: confirm each by reading the pages.", ""]
    if not report["findings"]:
        lines.append("No cross-page drift found.")
    for finding in report["findings"]:
        pages = ", ".join(f"P{p:02d}" for p in finding["pages"])
        lines += [f"## {finding['id']} {finding['severity'].upper()} {finding['kind']} - {pages}", "", finding["message"], "",
                  f"Fix: {finding['fix']}", ""]
        for element in finding["elements"]:
            lines.append(f"- P{element['page']:02d} `{element['element']}`: {element['excerpt']}")
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def repair_items(report: dict, certain_only: bool = True) -> dict[int, list[str]]:
    """Page-addressed repair hints: {page number: [text]}. Certain findings by default."""
    out: dict[int, list[str]] = defaultdict(list)
    for finding in report.get("findings") or []:
        if certain_only and finding["severity"] != CERTAIN:
            continue
        for page, text in (finding.get("repair") or {}).items():
            out[int(page)].append(f"[{finding['id']} {finding['kind']}] {text}")
    return dict(out)


def prompt_block(report: dict, limit: int = 40) -> str:
    """The findings as a block for a reviewer prompt: deterministic items to confirm or reject from the pages."""
    if not report.get("findings"):
        return "DETERMINISTIC CROSS-PAGE CHECK: nothing found."
    lines = [f"DETERMINISTIC CROSS-PAGE FINDINGS TO CONFIRM ({report['certain']} certain, {report['flagged']} flagged). A script read the drawn "
             "text of every page and found these. Confirm or reject each by reading the pages; do not report a figure the transcript does not contain."]
    for finding in report["findings"][:limit]:
        pages = ", ".join(f"P{p:02d}" for p in finding["pages"])
        lines.append(f"- {finding['id']} {finding['severity'].upper()} {finding['kind']} ({pages}): {finding['message']} Suggested: {finding['fix']}")
    return "\n".join(lines)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("target", help="project folder, folder of page SVGs, or an exported .pptx")
    parser.add_argument("--spec", help="design_spec.md (default: the project's own)")
    parser.add_argument("--json", dest="json_out", help="write the JSON report here")
    parser.add_argument("--md", dest="md_out", help="write the markdown report here")
    parser.add_argument("--print-json", action="store_true", help="print JSON instead of markdown")
    parser.add_argument("--fail-on-certain", action="store_true")
    args = parser.parse_args()
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")
    target = Path(args.target).resolve()
    spec = Path(args.spec).resolve() if args.spec else (target / "design_spec.md" if target.is_dir() else None)
    report = check(target, spec)
    markdown = render_markdown(report, target.name)
    is_project = target.is_dir() and (target / "svg_output").is_dir()
    json_out = Path(args.json_out) if args.json_out else (target / ".review" / "consistency.json" if is_project else None)
    md_out = Path(args.md_out) if args.md_out else (target / ".review" / "consistency.md" if is_project else None)
    for path, content in ((json_out, json.dumps(report, indent=2, ensure_ascii=False)), (md_out, markdown)):
        if path:
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(content, encoding="utf-8")
    print(json.dumps(report, indent=2, ensure_ascii=False) if args.print_json else markdown)
    return 1 if args.fail_on_certain and report["certain"] else 0


if __name__ == "__main__":
    sys.exit(main())
